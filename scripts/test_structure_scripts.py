"""Regression test for the structure-comments scripts (paragraph_index.py, add_comments.py).

Builds a synthetic draft in a temporary folder, runs both scripts as separate
processes the way the workflow does, and checks the documented behavior.
Nothing outside the temporary folder is written.

Run from any folder:

    python scripts/test_structure_scripts.py

Prints PASS or FAIL for each check (ASCII only); exits 0 when every check passes, 1 otherwise.
"""
import sys

sys.dont_write_bytecode = True

import hashlib  # noqa: E402
import json  # noqa: E402
import os  # noqa: E402
import shutil  # noqa: E402
import subprocess  # noqa: E402
import tempfile  # noqa: E402
import zipfile  # noqa: E402
from pathlib import Path  # noqa: E402

from docx import Document  # noqa: E402

SCRIPTS = Path(__file__).resolve().parent
ENV = dict(os.environ, PYTHONDONTWRITEBYTECODE="1", PYTHONIOENCODING="utf-8")
LQ, RQ, APOS = "\u201c", "\u201d", "\u2019"
results = []


def check(name, ok, detail=""):
    results.append(bool(ok))
    print("%s  %s" % ("PASS" if ok else "FAIL", name))
    if detail and not ok:
        print("        " + str(detail)[-600:].encode("ascii", "replace").decode())


def run(args):
    p = subprocess.run([sys.executable] + [str(a) for a in args], env=ENV,
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    return p.returncode, (p.stdout + p.stderr)


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def body_texts(path):
    from docx.oxml.ns import qn
    doc = Document(str(path))
    return ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))]


def comment_count(path):
    with zipfile.ZipFile(path) as z:
        if "word/comments.xml" not in z.namelist():
            return 0
        return z.read("word/comments.xml").decode("utf-8").count("<w:comment ")


def make_draft(path):
    doc = Document()
    doc.add_paragraph("Team autonomy and product recalls", style="Title")
    doc.add_paragraph("Abstract. We study how decision autonomy on plant teams relates to recall rates.")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Plants differ in how much authority their teams hold over the line.")
    doc.add_paragraph("Managers call this the " + LQ + "stop-the-line" + RQ + " right, and it"
                      + APOS + "s rarely written down.")
    doc.add_heading("Theory and hypotheses", level=1)
    doc.add_paragraph("Hypothesis 1: team decision autonomy lowers the recall rate.")
    doc.add_paragraph("The same sentence appears twice.")
    doc.add_paragraph("The same sentence appears twice.")
    doc.add_heading("Methods", level=1)
    doc.add_paragraph("We observed 412 plants from 2011 through 2019.")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Variable", "Mean"
    t.cell(1, 0).text, t.cell(1, 1).text = "Recall rate", "0.042"
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Table 2 reports the estimates for Hypothesis 1.")
    doc.add_heading("Discussion", level=1)
    doc.add_paragraph("Autonomy matters most where changeovers are frequent. It matters less on stable lines.")
    doc.save(str(path))


def comment_ranges(path):
    """{comment id: highlighted text} read from document.xml."""
    from lxml import etree
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    open_ids, got = set(), {}
    for el in root.iter():
        if el.tag == w + "commentRangeStart":
            open_ids.add(el.get(w + "id"))
            got.setdefault(el.get(w + "id"), "")
        elif el.tag == w + "commentRangeEnd":
            open_ids.discard(el.get(w + "id"))
        elif el.tag == w + "t":
            for cid in open_ids:
                got[cid] += el.text or ""
    return got


def main():
    with tempfile.TemporaryDirectory(prefix="structure_scripts_test_") as td:
        tmp = Path(td)
        draft = tmp / "draft.docx"
        make_draft(draft)
        before = sha(draft)
        texts_before = body_texts(draft)

        code, out = run([SCRIPTS / "paragraph_index.py", draft, "--out", tmp / "paragraph_index.json"])
        check("paragraph_index.py exits 0", code == 0, out)
        idx_path, map_path = tmp / "paragraph_index.json", tmp / "paragraph_index_map.txt"
        items = json.loads(idx_path.read_text(encoding="utf-8")) if idx_path.exists() else []
        check("index JSON written with one entry per paragraph", len(items) >= 15, "entries=%d" % len(items))
        check("index includes table-cell paragraphs", any(i.get("in_table") for i in items))
        check("index marks headings", any(i.get("is_heading") for i in items))
        map_ok = map_path.exists() and all(ord(ch) < 128 for ch in map_path.read_text(encoding="utf-8"))
        check("map file written and ASCII-safe", map_ok)

        comments = [
            {"anchor_quote": 'the "stop-the-line" right, and it\'s rarely', "comment": "First test comment."},
            {"anchor_quote": "We observed 412 plants", "comment": "Second test comment."},
        ]
        cjson = tmp / "comments_final.json"
        cjson.write_text(json.dumps(comments), encoding="utf-8")
        out1 = tmp / "draft_structure_comments.docx"
        code, out = run([SCRIPTS / "add_comments.py", draft, cjson, "--out", out1])
        check("add_comments.py injects two comments (exit 0)", code == 0 and "injected 2" in out, out)
        check("straight-quote anchor matched curly text", code == 0, out)
        check("the copy holds exactly two comments", out1.exists() and comment_count(out1) == 2,
              "count=%s" % (comment_count(out1) if out1.exists() else "no file"))
        check("body text of the copy equals the draft", out1.exists() and body_texts(out1) == texts_before)

        # Empty review results still use the normal validation and copy path.
        empty_json = tmp / "empty.json"
        empty_json.write_text("[]", encoding="utf-8")
        empty_out = tmp / "empty" / "draft_structure_comments.docx"
        code, out = run([SCRIPTS / "add_comments.py", draft, empty_json, "--out", empty_out])
        check("empty list succeeds and reports zero injected comments", code == 0 and "injected 0 comment(s)" in out, out)
        check("empty list creates a byte-identical copy with no new comments",
              empty_out.exists() and sha(empty_out) == before and comment_count(empty_out) == 0)

        empty_commented_out = tmp / "empty_commented.docx"
        commented_before = sha(out1)
        code, out = run([SCRIPTS / "add_comments.py", out1, empty_json, "--out", empty_commented_out])
        check("empty list preserves existing comments and the source bytes",
              code == 0 and empty_commented_out.exists()
              and sha(empty_commented_out) == commented_before and sha(out1) == commented_before, out)

        existing_empty_out = tmp / "existing_empty.docx"
        existing_empty_out.write_bytes(b"existing output must be preserved")
        existing_before = sha(existing_empty_out)
        code, out = run([SCRIPTS / "add_comments.py", draft, empty_json, "--out", existing_empty_out])
        check("empty list refuses existing output without --force and preserves it",
              code == 6 and sha(existing_empty_out) == existing_before, out)
        code, out = run([SCRIPTS / "add_comments.py", draft, empty_json, "--out", existing_empty_out, "--force"])
        check("empty list honors --force with a byte-identical replacement",
              code == 0 and "injected 0 comment(s)" in out and sha(existing_empty_out) == before, out)

        code, out = run([SCRIPTS / "add_comments.py", draft, empty_json, "--out", draft, "--force"])
        check("empty list cannot overwrite its input even with --force",
              code == 5 and sha(draft) == before, out)
        missing_input_out = tmp / "missing_input_out.docx"
        code, out = run([SCRIPTS / "add_comments.py", tmp / "absent.docx", empty_json, "--out", missing_input_out])
        check("empty list still rejects a missing input", code == 5 and not missing_input_out.exists(), out)
        invalid_draft = tmp / "invalid.docx"
        invalid_draft.write_bytes(b"not a DOCX package")
        invalid_before = sha(invalid_draft)
        invalid_out = tmp / "invalid_input_out.docx"
        code, out = run([SCRIPTS / "add_comments.py", invalid_draft, empty_json, "--out", invalid_out])
        check("empty list still parses the input DOCX before writing",
              code != 0 and not invalid_out.exists() and sha(invalid_draft) == invalid_before, out)

        for label, value in [("object", {}), ("null", None), ("string", "invalid"), ("number", 0)]:
            invalid_json = tmp / ("invalid_" + label + ".json")
            invalid_json.write_text(json.dumps(value), encoding="utf-8")
            invalid_json_out = tmp / ("invalid_" + label + "_out.docx")
            code, out = run([SCRIPTS / "add_comments.py", draft, invalid_json, "--out", invalid_json_out])
            check("non-list JSON " + label + " refused (exit 5), nothing written",
                  code == 5 and not invalid_json_out.exists(), out)

        missing = [{"anchor_quote": "a phrase that is not in the draft", "comment": "x"}]
        mjson = tmp / "missing.json"
        mjson.write_text(json.dumps(missing), encoding="utf-8")
        out2 = tmp / "missing_out.docx"
        code, out = run([SCRIPTS / "add_comments.py", draft, mjson, "--out", out2])
        check("missing anchor refused (exit 2), nothing written", code == 2 and not out2.exists(), out)

        dup = [{"anchor_quote": "The same sentence appears twice", "comment": "x"}]
        djson = tmp / "dup.json"
        djson.write_text(json.dumps(dup), encoding="utf-8")
        out3 = tmp / "dup_out.docx"
        code, out = run([SCRIPTS / "add_comments.py", draft, djson, "--out", out3])
        check("ambiguous anchor refused (exit 2), nothing written", code == 2 and not out3.exists(), out)

        code, out = run([SCRIPTS / "add_comments.py", draft, cjson, "--out", out1])
        check("existing output refused without --force (exit 6)", code == 6, out)
        code, out = run([SCRIPTS / "add_comments.py", draft, cjson, "--out", out1, "--force"])
        check("--force rewrites the output (exit 0)", code == 0, out)

        code, out = run([SCRIPTS / "add_comments.py", draft, cjson, "--out", draft])
        check("output path equal to the draft refused (exit 5)", code == 5, out)

        sentence = [{"anchor_quote": "matters less on stable lines", "comment": "Sentence test."}]
        sjson = tmp / "sentence.json"
        sjson.write_text(json.dumps(sentence), encoding="utf-8")
        out4 = tmp / "sentence_out.docx"
        code, out = run([SCRIPTS / "add_comments.py", draft, sjson, "--out", out4])
        got = comment_ranges(out4) if out4.exists() else {}
        check("a comment highlights only the sentence that holds its quote",
              code == 0 and list(got.values()) == ["It matters less on stable lines."]
              and "1 on a sentence" in out, (got, out))
        check("the sentence comment leaves the body text unchanged", out4.exists() and body_texts(out4) == texts_before)

        from docx.oxml import parse_xml
        box_draft = tmp / "box_draft.docx"
        doc = Document()
        p = doc.add_paragraph("A host paragraph that anchors a floating text box. ")
        p._p.append(parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape style="position:absolute;width:120pt;'
            'height:40pt"><v:textbox><w:txbxContent><w:p><w:r><w:t>Words that sit inside the box</w:t></w:r>'
            '</w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>'))
        doc.save(str(box_draft))
        bjson = tmp / "box.json"
        bjson.write_text(json.dumps([{"anchor_quote": "Words that sit inside the box", "comment": "x"}]),
                         encoding="utf-8")
        out5 = tmp / "box_out.docx"
        code, out = run([SCRIPTS / "add_comments.py", box_draft, bjson, "--out", out5])
        check("a quote taken from a text box is refused as ambiguous (exit 2), nothing written",
              code == 2 and "2 paragraphs" in out and not out5.exists(), out)

        # --- footnote and endnote texts reach the index in their own field ---
        from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part
        wns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

        def notes_part(kind, notes):
            seps = ('<w:{k} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{k}>'
                    '<w:{k} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/>'
                    '</w:r></w:p></w:{k}>').format(k=kind)
            body = "".join('<w:{k} w:id="{i}"><w:p><w:r><w:{k}Ref/></w:r><w:r><w:t xml:space="preserve"> {t}'
                           '</w:t></w:r></w:p></w:{k}>'.format(k=kind, i=i, t=t) for i, t in notes)
            return ('<w:{k}s xmlns:w="{ns}">{s}{b}</w:{k}s>'.format(k=kind, ns=wns, s=seps, b=body)).encode("utf-8")

        note_draft = tmp / "note_draft.docx"
        doc = Document()
        p = doc.add_paragraph("We coded autonomy from plant handbooks.")
        p._p.append(parse_xml('<w:r xmlns:w="%s"><w:footnoteReference w:id="1"/></w:r>' % wns))
        p = doc.add_paragraph("Recall notices came from a public registry.")
        p._p.append(parse_xml('<w:r xmlns:w="%s"><w:endnoteReference w:id="1"/></w:r>' % wns))
        doc.add_paragraph("A paragraph with no notes.")
        doc.part.relate_to(Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, notes_part(
            "footnote", [("1", "Two coders read every handbook; agreement was 0.91.")]), doc.part.package), RT.FOOTNOTES)
        doc.part.relate_to(Part(PackURI("/word/endnotes.xml"), CT.WML_ENDNOTES, notes_part(
            "endnote", [("1", "The registry lists every recall since 2005.")]), doc.part.package), RT.ENDNOTES)
        doc.save(str(note_draft))
        code, out = run([SCRIPTS / "paragraph_index.py", note_draft, "--out", tmp / "note_index.json"])
        nitems = json.loads((tmp / "note_index.json").read_text(encoding="utf-8")) if code == 0 else []
        check("footnote and endnote texts reach the index in a notes field; the paragraph text is unchanged",
              len(nitems) == 3 and nitems[0].get("text") == "We coded autonomy from plant handbooks."
              and nitems[0].get("notes") == [{"kind": "footnote", "id": "1",
                                              "text": "Two coders read every handbook; agreement was 0.91."}]
              and nitems[1].get("notes") == [{"kind": "endnote", "id": "1",
                                              "text": "The registry lists every recall since 2005."}]
              and nitems[2].get("notes") == [], (out, nitems))
        nmap = (tmp / "note_index_map.txt").read_text(encoding="utf-8") if code == 0 else ""
        check("the map marks the paragraphs that carry notes",
              "[notes: footnote 1]" in nmap and "[notes: endnote 1]" in nmap, nmap)

        # a missing note, and a note of two paragraphs
        note_draft2 = tmp / "note_draft2.docx"
        doc = Document()
        p = doc.add_paragraph("A mark whose note is missing.")
        p._p.append(parse_xml('<w:r xmlns:w="%s"><w:footnoteReference w:id="9"/></w:r>' % wns))
        p = doc.add_paragraph("A mark whose note has two paragraphs.")
        p._p.append(parse_xml('<w:r xmlns:w="%s"><w:footnoteReference w:id="2"/></w:r>' % wns))
        two_paras = notes_part("footnote", [("2", "First part.")]).replace(
            b"First part.</w:t></w:r></w:p>", b"First part.</w:t></w:r></w:p><w:p><w:r><w:t>Second part.</w:t></w:r></w:p>")
        doc.part.relate_to(Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, two_paras, doc.part.package),
                           RT.FOOTNOTES)
        doc.save(str(note_draft2))
        code, out = run([SCRIPTS / "paragraph_index.py", note_draft2, "--out", tmp / "note_index2.json"])
        n2 = json.loads((tmp / "note_index2.json").read_text(encoding="utf-8")) if code == 0 else []
        check("a mark whose note is missing reads [note text not found]",
              len(n2) == 2 and n2[0]["notes"][0]["text"] == "[note text not found]", (out, n2))
        check("a note of two paragraphs reads as one line, joined with a space",
              len(n2) == 2 and n2[1]["notes"][0]["text"] == "First part. Second part.", (out, n2))

        # a quote that occurs twice inside its own paragraph is refused
        twice = tmp / "twice.docx"
        doc = Document()
        doc.add_paragraph("Teams act fast when they can. Teams act fast when they must.")
        doc.save(str(twice))
        tjson = tmp / "twice.json"
        tjson.write_text(json.dumps([{"anchor_quote": "Teams act fast when they", "comment": "x"}]), encoding="utf-8")
        code, out = run([SCRIPTS / "add_comments.py", twice, tjson, "--out", tmp / "twice_out.docx"])
        check("a quote found twice inside one paragraph is refused (exit 2), nothing written",
              code == 2 and "times in one paragraph" in out and not (tmp / "twice_out.docx").exists(), out)
        # a repeat that differs only in spacing is a repeat for the placement too, so it is refused as well
        spaced = tmp / "spaced.docx"
        doc = Document()
        doc.add_paragraph("Teams  act fast when they can. Teams act fast when they must.")
        doc.save(str(spaced))
        code, out = run([SCRIPTS / "add_comments.py", spaced, tjson, "--out", tmp / "spaced_out.docx"])
        check("a quote repeated with different spacing inside one paragraph is refused too (exit 2)",
              code == 2 and "times in one paragraph" in out and not (tmp / "spaced_out.docx").exists(), out)

        # --- Windows paths over 260 characters ---
        if os.name == "nt":
            prefix = chr(92) * 2 + "?" + chr(92)
            deep = tmp / ("d" * 90) / ("e" * 90) / ("f" * 60)
            os.makedirs(prefix + str(deep))
            plain = str(deep / "draft_with_a_long_file_name.docx")
            shutil.copyfile(str(draft), prefix + plain)
            code, out = run([SCRIPTS / "paragraph_index.py", plain, "--out", str(deep / "paragraph_index.json")])
            indexed = code == 0 and os.path.exists(prefix + str(deep / "paragraph_index.json"))
            Path(prefix + str(deep / "comments_final.json")).write_text(json.dumps(comments), encoding="utf-8")
            code2, out2 = run([SCRIPTS / "add_comments.py", plain, str(deep / "comments_final.json"),
                               "--out", str(deep / "draft_structure_comments.docx")])
            check("paragraph_index.py and add_comments.py work when paths pass 260 characters",
                  len(plain) > 260 and indexed and code2 == 0
                  and os.path.exists(prefix + str(deep / "draft_structure_comments.docx")), out + out2)
            shutil.rmtree(prefix + str(tmp / ("d" * 90)))

        check("draft byte-identical after all runs (SHA-256)", sha(draft) == before)

        sys.path.insert(0, str(SCRIPTS))
        from docx_utils import verify_docx_formatting  # the skill's own copy
        res = verify_docx_formatting(out1, draft)
        check("verify_docx_formatting passes on the commented copy", res.get("passed") and not res.get("errors"),
              res)
        loaded = Path(sys.modules["docx_utils"].__file__).resolve()
        check("verify_docx_formatting came from the skill's own scripts/docx_utils.py",
              loaded == (SCRIPTS / "docx_utils.py").resolve(), str(loaded))

    print("%d of %d checks passed" % (sum(results), len(results)))
    return 0 if all(results) else 1


if __name__ == "__main__":
    sys.exit(main())
