"""Index a manuscript .docx for proofreading.

First checks all Word XML parts for unresolved tracked changes: if there are any, it
stops with exit code 3 and writes nothing (the author accepts or rejects them
in Word first).

Then writes five files next to each other in the output folder:
  para_index.txt         P### [style] text, one line per body paragraph; each
                         footnote or endnote mark appears as [^N] where it sits
  locked_paragraphs.txt  paragraphs the program will not rewrite, with the reason
                         (docx_utils.paragraph_lock_reason: a content control such
                         as a reference-manager citation, a simple field, an
                         equation, a symbol-font character, superscript or
                         subscript text, a tab or manual line break, a picture
                         placed between words, formatting that changes what the
                         text shows, ...)
  tables.txt             every Word table rendered as pipe-separated rows
  sections.txt           which standard sections look present, with evidence
  notes.txt              the text of each footnote and endnote, by paragraph and
                         [^N] mark, for reading only (no note is ever edited)

Why tables: python-docx's doc.paragraphs excludes table cells, so a
section agent that is told to check the tables cannot see them unless
they are extracted separately.

Why sections: a file may hold only part of a manuscript (one study, or one
section). Chunking needs to know which sections are present, so no agent is
dispatched for a section the file does not contain.

Usage:
  python index_manuscript.py --input draft.docx --out-dir language_editor_output/draft
"""
import sys

sys.dont_write_bytecode = True  # never leave a __pycache__ folder inside the skill

import argparse  # noqa: E402
import re  # noqa: E402
from pathlib import Path  # noqa: E402

from docx import Document  # noqa: E402

from editor_io import load_docx_utils  # noqa: E402

SECTION_CUES = [
    ("abstract", r"^\s*abstract\b"),
    ("introduction", r"^\s*(introduction|1\.?\s+introduction)\s*$"),
    ("theory", r"^\s*(theory|theoretical (background|framework|development)|literature review|"
                 r"hypothes[ie]s( development)?|theory and hypotheses|background)\b"),
    ("methods", r"^\s*(methods?|methodology|data and methods?|research (design|setting|methods?)|"
                  r"empirical (setting|strategy|approach)|study \d+ method)\b"),
    ("results", r"^\s*(results?|findings?|analys[ie]s and results?|empirical results?)\b"),
    ("discussion", r"^\s*(discussion|general discussion|discussion and conclusion|conclusions?|"
                     r"theoretical (implications|contributions))\b"),
]


def render_table(tbl):
    rows = []
    for row in tbl.rows:
        cells = [re.sub(r"\s+", " ", c.text).strip() for c in row.cells]
        rows.append(" | ".join(cells))
    return rows


def looks_like_heading(para):
    style = (para.style.name or "") if para.style is not None else ""
    if style.lower().startswith("heading") or style.lower() in ("title", "subtitle"):
        return True
    text = para.text.strip()
    if not text or len(text) > 80:
        return False
    # Short, no terminal period, mostly title case or all caps: a heading in a
    # draft that uses Normal style throughout (common in working manuscripts).
    return not text.endswith(".") and (text.isupper() or text == text.title() or len(text.split()) <= 6)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", required=True)
    ap.add_argument("--out-dir", required=True)
    args = ap.parse_args(argv)

    du = load_docx_utils(verbose=False)
    src = du.long_path(args.input)          # paths over 260 characters work on Windows
    outdir = du.long_path(args.out_dir)
    doc = Document(str(src))

    # --- unresolved tracked changes: stop before anything is written ---
    from docx_revisions import read_package, unresolved_revisions
    found_revisions = unresolved_revisions(read_package(src))
    tracked = len(found_revisions)
    if tracked:
        print("STOP: the draft has %d unresolved tracked change(s) (insertions, deletions, moves or "
              "formatting changes). Accept or reject all tracked changes in Word (Review, then Accept "
              "or Reject, then All Changes), save the draft, and run again. Nothing was written." % tracked)
        print("Affected Word parts: " + ", ".join(sorted({r["part"] for r in found_revisions})))
        return 3
    outdir.mkdir(parents=True, exist_ok=True)

    # --- paragraph index (the editors' view: footnote and endnote marks as [^N]) ---
    lines, locked, nonempty, words = [], [], 0, 0
    for i, p in enumerate(doc.paragraphs):
        style = (p.style.name if p.style is not None else "?") or "?"
        # ASCII white space only: a non-breaking space stays, so an unchanged
        # sentence copied into REVISED keeps it
        text = re.sub(r"[ \t\r\n\f\v]+", " ", du.paragraph_view_text(p._p)).strip(" ")
        lines.append("P%03d [%s] %s" % (i, style, text))
        if text:
            nonempty += 1
            words += len(text.split())
            reason = du.paragraph_lock_reason(p._p)
            if reason:
                locked.append("P%03d holds %s" % (i, reason))
    para_path = du.long_path(outdir / "para_index.txt")
    para_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    # --- paragraphs the program will not rewrite (flags and comments are still fine) ---
    locked_path = du.long_path(outdir / "locked_paragraphs.txt")
    locked_path.write_text(
        "Paragraphs the program will not rewrite, with the reason. Propose no REVISE for them;\n"
        "a flag (KEEP with a COMMENT) is fine.\n\n" + ("\n".join(locked) if locked else "(none)") + "\n",
        encoding="utf-8")

    # --- tables ---
    tbl_lines = []
    for n, tbl in enumerate(doc.tables, start=1):
        rows = render_table(tbl)
        tbl_lines.append("=== TABLE %d: %d rows x %d cols ===" % (n, len(tbl.rows), len(tbl.columns)))
        tbl_lines.extend(rows)
        tbl_lines.append("")
    tables_path = du.long_path(outdir / "tables.txt")
    tables_path.write_text("\n".join(tbl_lines) + "\n" if tbl_lines else "(no tables in this document)\n",
                           encoding="utf-8")

    # --- footnote and endnote texts, so an editor reads a note before flagging something as missing ---
    texts = du.note_texts(doc)
    note_lines = []
    for i, p in enumerate(doc.paragraphs):
        for n in du.paragraph_notes(p._p, texts):
            note_lines.append("P%03d [^%s] %s: %s" % (i, n["id"], n["kind"], n["text"]))
    notes_path = du.long_path(outdir / "notes.txt")
    notes_path.write_text(
        "Footnote and endnote texts, by paragraph and mark. For reading only: no note is edited.\n"
        "Read a paragraph's notes before flagging something in it as missing.\n\n"
        + ("\n".join(note_lines) if note_lines else "(no footnotes or endnotes)") + "\n", encoding="utf-8")

    # --- section presence ---
    found = {}
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue
        low = text.lower()
        if "abstract" not in found and re.match(SECTION_CUES[0][1], low):
            found["abstract"] = (i, text)  # a run-in "Abstract." label opens a paragraph of any length
            continue
        if len(text) > 120:
            continue
        for key, pattern in SECTION_CUES:
            if key in found:
                continue
            if re.match(pattern, low) and (looks_like_heading(p) or low.startswith("abstract")):
                found[key] = (i, text)
    sec_lines = ["Document: %s" % src.name,
                 "Body paragraphs: %d (%d non-empty), words: %d, tables: %d"
                 % (len(doc.paragraphs), nonempty, words, len(doc.tables)),
                 "",
                 "SECTION PRESENCE (heuristic: headings and cue words; confirm by reading)"]
    for key, _ in SECTION_CUES:
        if key in found:
            idx, text = found[key]
            sec_lines.append("  PRESENT  %-16s P%03d  %s" % (key, idx, text[:60]))
        else:
            sec_lines.append("  ABSENT   %-16s no heading or cue line matched" % key)
    missing = [k for k, _ in SECTION_CUES if k not in found]
    if missing:
        sec_lines += ["",
                      "No heading found for: " + ", ".join(missing) + ".",
                      "Chunk only the sections present; content under another heading is chunked there."]
    sections_path = du.long_path(outdir / "sections.txt")
    sections_path.write_text("\n".join(sec_lines) + "\n", encoding="utf-8")

    print("wrote %s (%d paragraphs)" % (para_path.name, len(doc.paragraphs)))
    print("wrote %s (%d locked)" % (locked_path.name, len(locked)))
    print("wrote %s (%d tables)" % (tables_path.name, len(doc.tables)))
    print("wrote %s" % sections_path.name)
    print("wrote %s (%d notes)" % (notes_path.name, len(note_lines)))
    for line in sec_lines[3:]:
        print(line.encode("ascii", "replace").decode("ascii"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
