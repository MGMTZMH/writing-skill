"""Portable tracked-output integration and preservation tests.

All fixtures are created in a temporary directory. No manuscript is modified.
Run with the Python interpreter used for requirements.txt. No Word is required.
"""
import sys
sys.dont_write_bytecode = True
import copy
import json
import os
import subprocess
import tempfile
import unittest
import zipfile
from types import SimpleNamespace
from io import BytesIO
import base64
from pathlib import Path
from unittest.mock import patch
from docx import Document
from lxml import etree
import docx_revisions as dr
import apply_language_edits as coordinator

HERE = Path(__file__).resolve().parent

def independent_text(root, accept):
    """Read revisions directly; does not call the writer's projection function."""
    def visit(node):
        if node.tag == dr.w("del") and accept or node.tag == dr.w("ins") and not accept:
            return ""
        if node.tag in {dr.w("t"), dr.w("delText")}:
            return node.text or ""
        if node.tag == dr.w("tab"):
            return "\t"
        if node.tag in {dr.w("br"), dr.w("cr")}:
            return "\n"
        return "".join(visit(c) for c in node)
    return [visit(p) for p in dr.body_paragraphs(root)]

class TrackedOutputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="tracked-tests-")
        self.folder = Path(self.temp.name)
        self.source = self.folder / "原稿 with spaces.docx"

    def tearDown(self):
        self.temp.cleanup()

    def draft(self, texts):
        doc = Document()
        for text in texts:
            doc.add_paragraph(text)
        doc.save(self.source)
        return doc

    def entry(self, n, text, revised=None, comment=None, anchor=None):
        return "\n".join(["### P%03d" % n, "ORIGINAL: " + text[:80],
                          "ACTION: " + ("REVISE" if revised is not None else "KEEP"),
                          "REVISED: " + revised if revised is not None else "COMMENT: null",
                          "COMMENT: " + comment if comment else "COMMENT: null",
                          "ANCHOR: " + anchor if anchor else "",
                          "IMPORTANCE: MINOR", "---", ""])

    def run_cli(self, entries, overrides=None, suffix="run", extra=None):
        folder = self.folder / suffix
        folder.mkdir(exist_ok=True)
        edits, ov = folder / "edits.md", folder / "overrides.json"
        edits.write_text("\n".join(entries), encoding="utf-8")
        ov.write_text(json.dumps(overrides or {}, ensure_ascii=False), encoding="utf-8")
        args = [sys.executable, str(HERE / "apply_language_edits.py"), "--input", str(self.source),
                "--out", str(folder / "clean.docx"), "--tracked-out", str(folder / "tracked.docx"),
                "--plan-out", str(folder / "plan.json"), "--manifest", str(folder / "manifest.json"),
                "--edits", str(edits), "--overrides", str(ov)] + (extra or [])
        env = dict(os.environ, PYTHONDONTWRITEBYTECODE="1", PYTHONIOENCODING="utf-8")
        env.pop("LANGUAGE_EDITOR_DOCX_UTILS", None)
        p = subprocess.run(args, capture_output=True, text=True, encoding="utf-8", env=env)
        return p, folder

    def assert_ready(self, result):
        p, folder = result
        self.assertEqual(p.returncode, 0, p.stdout + p.stderr)
        manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
        plan = json.loads((folder / "plan.json").read_text(encoding="utf-8"))
        self.assertEqual(manifest["status"], "ready")
        for record in manifest["outputs"].values():
            self.assertEqual(record["sha256"], dr.sha(Path(record["path"]).read_bytes()))
        root = dr.xml(dr.read_package(folder / "tracked.docx")[dr.DOCUMENT])
        original = dr.xml(dr.read_package(self.source)[dr.DOCUMENT])
        self.assertEqual(independent_text(root, False), independent_text(original, False))
        self.assertEqual(independent_text(root, True), [p.text for p in Document(folder / "clean.docx").paragraphs])
        return folder, manifest, plan, root

    def test_insert_delete_replace_noop_long_unicode(self):
        old = ["The authors report a relationship in their manuscript.",
               "The authors clearly report a relationship in their manuscript.",
               "The authors report a weak relationship in their manuscript.",
               "The unchanged sentence remains clear and accurate.",
               ("Long manuscript sentence with evidence. " * 20) + "中文 café ends."]
        new = [old[0].replace("report", "clearly report"), old[1].replace("clearly ", ""),
               old[2].replace("weak", "positive"), old[3], old[4].replace("café", "résumé")]
        self.draft(old)
        folder, m, plan, root = self.assert_ready(self.run_cli([self.entry(i,t,new[i]) for i,t in enumerate(old)]))
        self.assertEqual([r["approved_text"] for r in plan["paragraphs"]], new)
        self.assertGreater(len(plan["paragraphs"][4]["original_text"]), 400)
        self.assertEqual(plan["paragraphs"][3]["status"], "no_op")
        self.assertEqual(m["counts"]["no_op"], 1)
        self.assertEqual(sum(m["counts"][key] for key in ("applied", "reverted", "skipped", "no_op")), m["counts"]["proposed"])
        self.assertTrue(list(root.iter(dr.w("ins"))))
        self.assertTrue(list(root.iter(dr.w("delText"))))

    def test_full_and_partial_reversals(self):
        old = ["The finding may apply to smaller firms in this sample.", "The authors discuss two reasons for their choice."]
        self.draft(old)
        revised = [old[0].replace("may", "will"), old[1].replace("discuss two", "fully explain three")]
        approved = old[1].replace("discuss", "explain")
        _, m, plan, root = self.assert_ready(self.run_cli([self.entry(i,t,revised[i]) for i,t in enumerate(old)],
             {"P000":{"class":"reversal", "reason":"certainty"}, "P001":{"class":"partial", "revised":approved, "reason":"count"}}))
        self.assertEqual(plan["paragraphs"][0]["status"], "reverted")
        self.assertFalse(list(dr.body_paragraphs(root)[0].iter(dr.w("ins"))))
        self.assertEqual(plan["paragraphs"][1]["approved_text"], approved)
        self.assertEqual(m["counts"]["partial"], 1)

    def test_comments_final_sentences_and_drift(self):
        old = ["First sentence gives background. Second sentence needs a clearer explanation.",
               "The unchanged manuscript sentence needs missing facts."]
        self.draft(old)
        revised = old[0].replace("clearer explanation", "specific explanation")
        _, m, plan, root = self.assert_ready(self.run_cli([
            self.entry(0, old[0], revised, "Clarify the evidence and the scope together.", "specific explanation"),
            self.entry(1, old[1], comment="What population and time period does this refer to?", anchor="unchanged manuscript")],
            {"P000":{"class":"drift","notes":[{"anchor":"Second sentence", "note":"The revision makes the explanation more specific."}]}}))
        self.assertEqual(len(plan["comments"]), 3)
        self.assertEqual(plan["comments"][0]["anchor_text"], "Second sentence needs a specific explanation.")
        self.assertEqual(m["counts"]["comments_injected"], 3)

    def test_repeated_anchor_blocks(self):
        text = "Repeated word in first sentence. Repeated word in second sentence."
        self.draft([text])
        p, folder = self.run_cli([self.entry(0,text,comment="Clarify.",anchor="Repeated word")])
        self.assertEqual(p.returncode, 4, p.stdout+p.stderr)
        self.assertFalse((folder / "tracked.docx").exists())

    def test_deleted_anchor_blocks(self):
        text = "The manuscript has an unclear sentence about the argument."
        self.draft([text])
        p, folder = self.run_cli([self.entry(0,text,text.replace("unclear", "clear"), "Clarify.", "unclear")])
        self.assertEqual(p.returncode, 4)
        self.assertFalse((folder / "tracked.docx").exists())

    def test_no_anchor_blocks(self):
        text = "This sentence in the original needs an author decision."
        self.draft([text])
        p, folder = self.run_cli([self.entry(0,text,comment="Clarify.")])
        self.assertEqual(p.returncode, 4)

    def test_mixed_formatting_retained_and_ambiguous_skip(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("This plain sentence retains ")
        p.add_run("bold emphasis").bold = True
        p.add_run(" and its meaningful ending.")
        doc.save(self.source)
        text = p.text
        _, m, plan, root = self.assert_ready(self.run_cli([self.entry(0,text,text.replace("plain", "clear"))]))
        self.assertEqual(m["counts"]["applied"], 1)
        self.assertEqual([r.text for r in Document(self.folder/'run'/'clean.docx').paragraphs[0].runs if r.bold], ["bold emphasis"])
        _, m, plan, root = self.assert_ready(self.run_cli([self.entry(0,text,text.replace("retains bold emphasis", "preserves emphasis"))],suffix="skip"))
        self.assertEqual(m["counts"]["skipped"], 1)

    def test_italic_replaced_span_inherits_style(self):
        doc = Document()
        p = doc.add_paragraph("The manuscript describes an ")
        p.add_run("unclear").italic = True
        p.add_run(" relationship between the variables.")
        doc.save(self.source)
        folder, _, _, _ = self.assert_ready(self.run_cli([self.entry(0,p.text,p.text.replace("unclear","inverse"))]))
        self.assertTrue(any(r.text == "inverse" and r.italic for r in Document(folder/'clean.docx').paragraphs[0].runs))

    def test_bookmarks_and_existing_comments_preserved(self):
        doc = self.draft(["An initial sentence contains a bookmark. The second sentence needs revision."])
        p = doc.paragraphs[0]
        p._p.insert(0, etree.Element(dr.w("bookmarkStart"), {dr.w("id"):"42",dr.w("name"):"sample"}))
        p._p.append(etree.Element(dr.w("bookmarkEnd"), {dr.w("id"):"42"}))
        doc.save(self.source)
        dr.du.inject_comments_ex(self.source,[{"para_index":0,"text":"Original comment","author":"Author","highlight":"initial sentence"}])
        text = p.text
        _, m, plan, root = self.assert_ready(self.run_cli([self.entry(0,text,text.replace("needs revision","needs clarification"))]))
        self.assertEqual(len(list(root.iter(dr.w("bookmarkStart")))),1)
        self.assertIn(m["counts"]["applied"], (0,1))

    def test_fields_hyperlinks_notes_tables_preserved(self):
        doc = Document()
        p = doc.add_paragraph("The authors clearly explain ")
        link = etree.SubElement(p._p,dr.w("hyperlink"),{dr.w("anchor"):"target"})
        etree.SubElement(etree.SubElement(link,dr.w("r")),dr.w("t")).text="the sample"
        p.add_run(" and cite ")
        for kind,tag,text in [('begin','fldChar',None),('', 'instrText',' CITATION Smith '),('separate','fldChar',None),('','t','(Smith, 2020)'),('end','fldChar',None)]:
            r=etree.SubElement(p._p,dr.w('r'));n=etree.SubElement(r,dr.w(tag))
            if kind:n.set(dr.w('fldCharType'),kind)
            if text:n.text=text
        p.add_run(" in this passage.")
        doc.add_table(rows=1,cols=1).cell(0,0).text="Protected table value 42"
        doc.save(self.source)
        text=dr.du.paragraph_view_text(p._p)
        _, m, plan, root=self.assert_ready(self.run_cli([self.entry(0,text,text.replace('clearly explain','describe'))]))
        self.assertEqual(m['counts']['applied'],1)
        self.assertEqual(len(list(root.iter(dr.w('hyperlink')))),1)
        self.assertEqual(len(list(root.iter(dr.w('fldChar')))),3)
        _,m,_,_=self.assert_ready(self.run_cli([self.entry(0,text,text.replace('the sample','the population'))],suffix='protected'))
        self.assertEqual(m['counts']['skipped'],1)

    def test_note_id_change_and_equation_tab_are_skipped(self):
        doc=Document();p=doc.add_paragraph('The original sentence has a note')
        ref=etree.SubElement(etree.SubElement(p._p,dr.w('r')),dr.w('footnoteReference'));ref.set(dr.w('id'),'1')
        p.add_run(' and its meaning stays clear.')
        q=doc.add_paragraph('A manuscript sentence has\ta tab in its middle.')
        r=doc.add_paragraph('This manuscript paragraph contains an equation. ')
        etree.SubElement(r._p,'{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')
        doc.save(self.source)
        old=[dr.du.paragraph_view_text(x._p) for x in doc.paragraphs]
        _,m,_,_=self.assert_ready(self.run_cli([self.entry(i,t,t.replace('original','revised').replace('[^1]','[^2]').replace('manuscript','clear')) for i,t in enumerate(old)]))
        self.assertEqual(m['counts']['skipped'],3)

    def test_repeated_identical_note_ids_have_distinct_objects(self):
        doc=Document();p=doc.add_paragraph('The initial sentence uses note')
        for kind in ('footnoteReference','endnoteReference'):
            ref=etree.SubElement(etree.SubElement(p._p,dr.w('r')),dr.w(kind));ref.set(dr.w('id'),'1')
            p.add_run(' and another note' if kind=='footnoteReference' else ' in its argument.')
        doc.save(self.source)
        old=dr.du.paragraph_view_text(p._p)
        _,m,_,root=self.assert_ready(self.run_cli([self.entry(0,old,old.replace('initial','first'))]))
        self.assertEqual(m['counts']['applied'],1)
        self.assertEqual(len(list(root.iter(dr.w('footnoteReference')))),1)
        self.assertEqual(len(list(root.iter(dr.w('endnoteReference')))),1)

    def test_format_revision_in_header_blocks_input(self):
        doc=self.draft(['The original manuscript sentence is sufficiently long.'])
        h=doc.sections[0].header.paragraphs[0]
        r=h.add_run('Header');pr=r._r.get_or_add_rPr();etree.SubElement(pr,dr.w('rPrChange'),{dr.w('id'):'99'})
        doc.save(self.source)
        p,folder=self.run_cli([self.entry(0,doc.paragraphs[0].text)])
        self.assertEqual(p.returncode,3,p.stdout+p.stderr)
        self.assertFalse((folder/'tracked.docx').exists())
        indexed=self.folder/'indexed'
        check=subprocess.run([sys.executable,str(HERE/'index_manuscript.py'),'--input',str(self.source),'--out-dir',str(indexed)],capture_output=True,text=True)
        self.assertEqual(check.returncode,3,check.stdout+check.stderr)
        self.assertFalse(indexed.exists())

    def test_move_revision_in_footnote_part_detected(self):
        self.draft(['The original manuscript sentence is sufficiently long.'])
        source=dr.read_package(self.source)
        source['word/footnotes.xml']=('<w:footnotes xmlns:w="%s"><w:footnote w:id="1"><w:p><w:moveFromRangeStart w:id="7"/></w:p></w:footnote></w:footnotes>' % dr.W).encode()
        self.assertTrue(any(r['tag']=='moveFromRangeStart' for r in dr.unresolved_revisions(source)))

    def test_duplicate_and_ambiguous_targets_block(self):
        text='The repeated manuscript paragraph has the same complete wording.'
        self.draft([text,text])
        p,folder=self.run_cli([self.entry(9,text,text.replace('complete','exact'))])
        self.assertEqual(p.returncode,5,p.stdout+p.stderr)
        self.assertFalse((folder/'tracked.docx').exists())
        p,folder=self.run_cli([self.entry(0,text,text.replace('complete','exact'))]*2,suffix='duplicate')
        self.assertNotEqual(p.returncode,0)

    def test_output_collision_does_not_overwrite(self):
        text='The original manuscript sentence contains enough words.'
        self.draft([text]);before=self.source.read_bytes()
        p,folder=self.run_cli([self.entry(0,text)],extra=['--tracked-out',str(self.source)])
        self.assertNotEqual(p.returncode,0)
        self.assertEqual(before,self.source.read_bytes())

    def test_existing_output_not_overwritten(self):
        text='The original manuscript sentence contains enough words.'
        self.draft([text]);result=self.run_cli([self.entry(0,text)]);self.assert_ready(result)
        folder=result[1];before=(folder/'tracked.docx').read_bytes()
        p,_=self.run_cli([self.entry(0,text)])
        self.assertNotEqual(p.returncode,0)
        self.assertEqual(before,(folder/'tracked.docx').read_bytes())

    def test_dependency_failure(self):
        with patch.object(coordinator.importlib.metadata,'version',side_effect=coordinator.importlib.metadata.PackageNotFoundError):
            with self.assertRaisesRegex(ValueError,'missing dependency'):
                coordinator.dependency_check()

    def test_hash_and_invalid_plan_and_tampering(self):
        old='The original manuscript sentence states the predicted relationship.'
        self.draft([old]);folder,m,plan,root=self.assert_ready(self.run_cli([self.entry(0,old,old.replace('states','describes'))]))
        source=dr.read_package(self.source)
        with self.assertRaisesRegex(dr.RevisionError,'source hash'):
            dr.build_output(source,plan,'wrong')
        with self.assertRaises(dr.RevisionError):
            dr.build_output(source,{},dr.sha(self.source.read_bytes()))
        tampered=dr.read_package(folder/'tracked.docx')
        root=dr.xml(tampered[dr.DOCUMENT]);next(root.iter(dr.w('delText'))).text='WRONG'
        tampered[dr.DOCUMENT]=dr.serialize(root)
        with self.assertRaisesRegex(dr.RevisionError,'rejected view'):
            dr.validate_result(source,tampered,plan)

    def test_nonbody_part_and_format_tamper_detected(self):
        old='The original manuscript sentence states the predicted relationship.'
        self.draft([old]);folder,m,plan,root=self.assert_ready(self.run_cli([self.entry(0,old,old.replace('states','describes'))]))
        source=dr.read_package(self.source);output=dr.read_package(folder/'tracked.docx')
        output['word/styles.xml']=output['word/styles.xml']+b'\n'
        with self.assertRaisesRegex(dr.RevisionError,'package part differs'):
            dr.validate_result(source,output,plan)
        output=dr.read_package(folder/'tracked.docx');root=dr.xml(output[dr.DOCUMENT])
        run=next(root.iter(dr.w('r')));pr=etree.Element(dr.w('rPr'));etree.SubElement(pr,dr.w('b'));run.insert(0,pr)
        output[dr.DOCUMENT]=dr.serialize(root)
        with self.assertRaises(dr.RevisionError):dr.validate_result(source,output,plan)

    def test_failed_copy_removes_partial_destination(self):
        stage=self.folder/'stage';stage.write_bytes(b'contents');dest=self.folder/'dest'
        with patch.object(coordinator.os,'fsync',side_effect=OSError('disk failure')):
            with self.assertRaises(OSError):coordinator.publish_file(stage,dest)
        self.assertFalse(dest.exists())

    def test_significant_spaces_and_insertion_wrapper_sentence_split(self):
        old='  The manuscript describes an unclear relationship.  '
        new='  The manuscript describes the relationship. The new explanation needs evidence.  '
        self.draft([old])
        folder,m,plan,root=self.assert_ready(self.run_cli([self.entry(0,old,new,'Which evidence supports it?','new explanation')]))
        self.assertEqual(plan['paragraphs'][0]['approved_text'],new)
        self.assertEqual(plan['comments'][0]['anchor_text'],'The new explanation needs evidence.')

    def test_custom_xml_revision_and_conflict_markers(self):
        self.draft(['The original manuscript sentence contains sufficient detail.'])
        source=dr.read_package(self.source)
        for namespace,name in [(dr.W,'customXmlMoveToRangeStart'),('http://schemas.microsoft.com/office/word/2010/wordml','conflictIns')]:
            root=dr.xml(source[dr.DOCUMENT]);etree.SubElement(dr.body_paragraphs(root)[0],'{%s}%s'%(namespace,name))
            parts=dict(source);parts[dr.DOCUMENT]=dr.serialize(root)
            self.assertTrue(any(r['tag']==name for r in dr.unresolved_revisions(parts)))
            path=self.folder/(name+'.docx');dr.write_package(path,parts);dest=self.folder/(name+'-index')
            check=subprocess.run([sys.executable,str(HERE/'index_manuscript.py'),'--input',str(path),'--out-dir',str(dest)],capture_output=True,text=True)
            self.assertEqual(check.returncode,3,check.stdout+check.stderr)
            self.assertFalse(dest.exists())

    def test_drawing_resource_preserved_and_object_edit_skipped(self):
        doc=Document();p=doc.add_paragraph('The original manuscript sentence precedes an inline picture. ')
        png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')
        p.add_run().add_picture(BytesIO(png))
        q=doc.add_paragraph('The second manuscript paragraph is unaffected by the picture.')
        doc.save(self.source)
        folder,m,plan,root=self.assert_ready(self.run_cli([self.entry(0,p.text,p.text.replace('original','revised')),self.entry(1,q.text,q.text.replace('second','next'))]))
        self.assertEqual(m['counts']['skipped'],1)
        source=dr.read_package(self.source);output=dr.read_package(folder/'tracked.docx')
        self.assertTrue(any(n.startswith('word/media/') for n in source))
        for n in source:
            if n.startswith('word/media/'):self.assertEqual(source[n],output[n])

    def direct_args(self):
        text='The original manuscript sentence states a testable prediction.'
        self.draft([text]);edits=self.folder/'edits.md';edits.write_text(self.entry(0,text,text.replace('states','expresses')))
        ov=self.folder/'overrides.json';ov.write_text('{}')
        return SimpleNamespace(input=str(self.source),out=str(self.folder/'clean.docx'),tracked_out=str(self.folder/'tracked.docx'),plan_out=str(self.folder/'plan.json'),manifest=str(self.folder/'manifest.json'),edits=[str(edits)],overrides=str(ov),para_index=None,author='Tester',rule_set=None)

    def test_changed_source_hash_blocks_publication(self):
        args=self.direct_args();du=dr.du;original_verify=du.verify_docx_formatting
        def mutate_source(*a,**kw):
            result=original_verify(*a,**kw)
            self.source.write_bytes(self.source.read_bytes()+b'changed during validation')
            return result
        with patch.object(coordinator,'load_docx_utils',return_value=du),patch.object(du,'verify_docx_formatting',side_effect=mutate_source):
            with self.assertRaisesRegex(dr.RevisionError,'source hash changed'):coordinator.run(args)
        self.assertFalse(Path(args.tracked_out).exists())
        self.assertFalse(Path(args.manifest).exists())

    def test_mid_publication_failure_marks_incomplete(self):
        args=self.direct_args();real=coordinator.publish_file;calls=[]
        def fail_second(stage,dest):
            calls.append(dest)
            if len(calls)==2:raise OSError('simulated publication failure')
            return real(stage,dest)
        with patch.object(coordinator,'publish_file',side_effect=fail_second):
            result=coordinator.run(args)
        self.assertEqual(result,6)
        self.assertFalse(Path(args.out).exists())
        self.assertFalse(Path(args.tracked_out).exists())
        self.assertEqual(json.loads(Path(args.manifest).read_text())['status'],'incomplete')

    def test_inserted_style_tamper_detected(self):
        old='The original manuscript sentence states the predicted relationship.'
        self.draft([old]);folder,m,plan,root=self.assert_ready(self.run_cli([self.entry(0,old,old.replace('states','describes'))]))
        source=dr.read_package(self.source);output=dr.read_package(folder/'tracked.docx');root=dr.xml(output[dr.DOCUMENT])
        run=next(root.iter(dr.w('ins'))).find(dr.w('r'));pr=etree.Element(dr.w('rPr'));etree.SubElement(pr,dr.w('i'));run.insert(0,pr)
        output[dr.DOCUMENT]=dr.serialize(root)
        with self.assertRaisesRegex(dr.RevisionError,'inserted formatting'):dr.validate_result(source,output,plan)

    def test_dangling_relationship_is_rejected(self):
        self.draft(['The original manuscript sentence contains sufficient detail.'])
        source=dr.read_package(self.source);root=dr.xml(source[dr.DOCUMENT])
        etree.SubElement(dr.body_paragraphs(root)[0],dr.w('hyperlink'),{'{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id':'rIdNotPresent'})
        source[dr.DOCUMENT]=dr.serialize(root)
        with self.assertRaisesRegex(dr.RevisionError,'unresolved relationship'):dr.validate_package(source)

    def test_partial_to_original_is_counted_as_no_op(self):
        old = "The finding may apply to smaller firms in this sample."
        self.draft([old])
        proposed = old.replace("may", "will")
        _, manifest, plan, root = self.assert_ready(self.run_cli(
            [self.entry(0, old, proposed)],
            {"P000": {"class": "partial", "revised": old, "reason": "Restore the unsupported certainty."}}))
        self.assertEqual(plan["paragraphs"][0]["status"], "no_op")
        self.assertEqual(manifest["counts"]["no_op"], 1)
        self.assertEqual(manifest["counts"]["partial"], 1)
        self.assertEqual(sum(manifest["counts"][key] for key in ("applied", "reverted", "skipped", "no_op")), manifest["counts"]["proposed"])
        self.assertFalse(list(root.iter(dr.w("ins"), dr.w("del"))))

    def test_unplanned_comment_body_blocks_validation(self):
        old = "The original manuscript sentence states the predicted relationship."
        self.draft([old])
        folder, _, plan, _ = self.assert_ready(self.run_cli([
            self.entry(0, old, old.replace("states", "describes"), "Clarify the relationship.", "predicted relationship")]))
        source = dr.read_package(self.source)
        output = dr.read_package(folder / "tracked.docx")
        comments = dr.xml(output["word/comments.xml"])
        orphan = etree.SubElement(comments, dr.w("comment"), {dr.w("id"): "999"})
        etree.SubElement(etree.SubElement(etree.SubElement(orphan, dr.w("p")), dr.w("r")), dr.w("t")).text = "Unplanned annotation"
        output["word/comments.xml"] = dr.serialize(comments)
        with self.assertRaisesRegex(dr.RevisionError, "unplanned or missing comment body"):
            dr.validate_result(source, output, plan)

    def test_paragraph_movement_across_table_blocks_validation(self):
        doc = Document()
        p = doc.add_paragraph("The paragraph before the table introduces the numerical evidence.")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Protected table"
        doc.add_paragraph("The paragraph after the table interprets the numerical evidence.")
        doc.save(self.source)
        folder, _, plan, _ = self.assert_ready(self.run_cli([self.entry(0, p.text)]))
        source = dr.read_package(self.source)
        output = dr.read_package(folder / "tracked.docx")
        root = dr.xml(output[dr.DOCUMENT])
        body = root.find(dr.w("body"))
        first = body[0]
        body.remove(first)
        body.insert(1, first)
        output[dr.DOCUMENT] = dr.serialize(root)
        with self.assertRaisesRegex(dr.RevisionError, "non-body-paragraph content or section structure changed"):
            dr.validate_result(source, output, plan)

if __name__=='__main__':
    unittest.main(verbosity=2)
