"""
docx_utils.py -- Word document utilities for the writing skill (both functions).

The proofreading scripts (index_manuscript.py, apply_language_edits.py) and the
structure-comment script (add_comments.py) do all .docx text replacement, comment
injection and verification through these functions, so every run treats text the
same way.

Provides:
    - smartify_quotes(): Convert straight quotes to curly/smart quotes
    - paragraph_view_text(): The paragraph as the editors see it, footnote marks as [^N]
    - paragraph_lock_reason(): Why a paragraph cannot be rewritten safely (or None)
    - count_tracked_changes(): Unresolved tracked changes inside an element
    - replace_paragraph_text(): Replace a paragraph's text, keeping formatting, footnote
      marks, links, reference-manager fields and floating text boxes; all or nothing
    - sentence_spans(): Where the sentences of a paragraph's text begin and end
    - inject_comments_ex(): Add Word comments that highlight a sentence (or a paragraph)
    - inject_comments(): The same, for (search_text, comment_text, author[, highlight]) tuples
    - verify_docx_formatting(): Check output .docx for formatting errors

Usage:
    from docx_utils import replace_paragraph_text, inject_comments_ex, verify_docx_formatting
"""

import os
import re
import copy
import datetime
import zipfile
from pathlib import Path
from io import BytesIO
from lxml import etree

# ---------------------------------------------------------------------------
# XML namespace
# ---------------------------------------------------------------------------
WML = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSMAP = {"w": WML}
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
WPD = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
VML = "urn:schemas-microsoft-com:vml"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _w(tag: str) -> str:
    """Clark name of a WordprocessingML tag: _w("t") -> "{...main}t"."""
    return f"{{{WML}}}{tag}"


def long_path(path) -> Path:
    """An absolute path that still works on Windows beyond 260 characters.

    Windows refuses longer paths unless long-path support is switched on, which
    it often is not; a deep OneDrive folder plus the output folder names can pass
    the limit. Paths over 240 characters get the \\\\?\\ prefix, which lifts it;
    a network path (one that starts with two backslashes) gets its UNC form.
    The scripts pass every file they open or write through this function.
    """
    p = Path(path)
    if not p.is_absolute():
        p = Path.cwd() / p
    s = os.path.normpath(str(p))
    if os.name == "nt" and len(s) > 240 and not s.startswith("\\\\?\\"):
        if s.startswith("\\\\"):
            return Path("\\\\?\\UNC\\" + s[2:])
        return Path("\\\\?\\" + s)
    return Path(s)


# ---------------------------------------------------------------------------
# Curly quote conversion
# ---------------------------------------------------------------------------
def smartify_quotes(text: str) -> str:
    """Convert straight quotes to curly/smart quotes."""
    result = []
    for i, ch in enumerate(text):
        if ch == '"':
            if i == 0 or text[i - 1] in ' \t\n\r([':
                result.append('\u201C')
            else:
                result.append('\u201D')
        elif ch == "'":
            if i == 0 or text[i - 1] in ' \t\n\r([':
                result.append('\u2018')
            else:
                result.append('\u2019')
        else:
            result.append(ch)
    return ''.join(result)


# ---------------------------------------------------------------------------
# Run formatting helpers
# ---------------------------------------------------------------------------
def get_run_formatting(run_el) -> tuple:
    """Return (is_bold, is_italic) based on explicit run properties."""
    rPr = run_el.find("w:rPr", NSMAP)
    if rPr is None:
        return (False, False)

    is_bold = False
    b = rPr.find("w:b", NSMAP)
    if b is not None:
        val = b.get(f"{{{WML}}}val")
        if val is None or val not in ("0", "false"):
            is_bold = True

    is_italic = False
    i_el = rPr.find("w:i", NSMAP)
    if i_el is not None:
        val = i_el.get(f"{{{WML}}}val")
        if val is None or val not in ("0", "false"):
            is_italic = True

    return (is_bold, is_italic)


def is_run_formatted(run_el) -> bool:
    """Return True if the run has explicit bold or italic."""
    b, i = get_run_formatting(run_el)
    return b or i


def get_run_text(run_el) -> str:
    """Text of the run's own <w:t> children.

    Only direct children count. A run that anchors a text box also holds the
    box's paragraphs; their <w:t> belong to the box, never to this run.
    """
    return "".join((t.text or "") for t in run_el.findall("w:t", NSMAP))


def _run_has_footnote_ref(run_el) -> bool:
    """True for a run holding a footnote or endnote reference mark."""
    return (run_el.find("w:footnoteReference", NSMAP) is not None
            or run_el.find("w:endnoteReference", NSMAP) is not None)


def _run_has_comment_ref(run_el) -> bool:
    return run_el.find("w:commentReference", NSMAP) is not None


# ---------------------------------------------------------------------------
# Text manipulation on runs (own <w:t> children only; see get_run_text)
# ---------------------------------------------------------------------------
def _set_run_text(run_el, text: str) -> None:
    t_elements = run_el.findall("w:t", NSMAP)
    if t_elements:
        t_elements[0].text = text
        t_elements[0].set(XML_SPACE, "preserve")
        for t_el in t_elements[1:]:
            t_el.text = ""


def _clear_run_text(run_el) -> None:
    for t_el in run_el.findall("w:t", NSMAP):
        t_el.text = ""


def _plain_run_like(run_el, text: str):
    """A new run with the same run properties and one text element.

    Nothing else in the run (tabs, breaks, pictures, text boxes) is copied.
    """
    new_run = etree.Element(_w("r"))
    rpr = run_el.find("w:rPr", NSMAP)
    if rpr is not None:
        new_run.append(copy.deepcopy(rpr))
    t_el = etree.SubElement(new_run, _w("t"))
    t_el.text = text
    t_el.set(XML_SPACE, "preserve")
    return new_run


# ---------------------------------------------------------------------------
# Formatting manipulation on runs
# ---------------------------------------------------------------------------
def _apply_formatting_to_run(run_el, should_bold, should_italic):
    """Set or remove bold/italic formatting on a run element."""
    rPr = run_el.find("w:rPr", NSMAP)

    if should_bold or should_italic:
        if rPr is None:
            rPr = etree.SubElement(run_el, f"{{{WML}}}rPr")
            # rPr must be the first child of w:r
            run_el.remove(rPr)
            run_el.insert(0, rPr)

        # Bold
        b_el = rPr.find("w:b", NSMAP)
        if should_bold and b_el is None:
            etree.SubElement(rPr, f"{{{WML}}}b")
        elif not should_bold and b_el is not None:
            rPr.remove(b_el)

        # Italic
        i_el = rPr.find("w:i", NSMAP)
        if should_italic and i_el is None:
            etree.SubElement(rPr, f"{{{WML}}}i")
        elif not should_italic and i_el is not None:
            rPr.remove(i_el)
    else:
        # Ensure no bold/italic on this run
        if rPr is not None:
            for tag_name in ["w:b", "w:i"]:
                el = rPr.find(tag_name, NSMAP)
                if el is not None:
                    rPr.remove(el)


# ---------------------------------------------------------------------------
# Tracked changes, locked paragraphs, and the editors' view of a paragraph
# ---------------------------------------------------------------------------
TRACKED_CHANGE_TAGS = frozenset(_w(t) for t in (
    "ins", "del", "moveFrom", "moveTo", "rPrChange", "pPrChange", "sectPrChange",
    "tblPrChange", "tblPrExChange", "trPrChange", "tcPrChange", "tblGridChange",
    "cellIns", "cellDel", "cellMerge", "numberingChange"))

OBJECT_TAGS = frozenset((_w("drawing"), _w("pict"), _w("object"), f"{{{MC}}}AlternateContent"))

# Paragraph (or link) children that hold no visible text of their own.
_NO_TEXT_CHILDREN = frozenset(_w(t) for t in (
    "pPr", "rPr", "bookmarkStart", "bookmarkEnd", "proofErr", "commentRangeStart",
    "commentRangeEnd", "permStart", "permEnd", "moveFromRangeStart", "moveFromRangeEnd",
    "moveToRangeStart", "moveToRangeEnd"))

# Run children the rewrite knows how to keep. A run holding anything else (a
# symbol-font character, a positional tab, ...) locks its paragraph.
_RUN_CHILDREN_OK = frozenset(_w(t) for t in (
    "rPr", "t", "tab", "br", "cr", "noBreakHyphen", "softHyphen", "footnoteReference",
    "endnoteReference", "fldChar", "instrText", "lastRenderedPageBreak", "commentReference",
    "annotationRef")) | OBJECT_TAGS

MATH = "http://schemas.openxmlformats.org/officeDocument/2006/math"

# A footnote or endnote mark as the editors see it in the paragraph index.
_NOTE_MARK = re.compile(r"\[\^[^\]\s]*\]")

# Curly quotes folded to straight ones, one character for one, so positions keep.
_QUOTE_FOLD = {0x2018: "'", 0x2019: "'", 0x201C: '"', 0x201D: '"'}


def _fold_quotes(s: str) -> str:
    return s.translate(_QUOTE_FOLD)


def _local(tag) -> str:
    return tag.split("}")[-1] if isinstance(tag, str) else "?"


def count_tracked_changes(element) -> int:
    """Number of unresolved tracked changes inside an element (a paragraph or the body).

    Counts insertions, deletions, moves, and formatting, section, table and
    numbering changes.
    """
    return sum(1 for el in element.iter() if el.tag in TRACKED_CHANGE_TAGS)


def _run_objects(run_el) -> list:
    """Drawings, pictures, text boxes and embedded objects held directly by a run."""
    return [child for child in run_el if child.tag in OBJECT_TAGS]


def _object_floats(obj) -> bool:
    """True when an object floats: anchored to the paragraph, positioned on the page.

    A floating object sits between no two words, so the text around it can be
    rewritten. An inline object sits between two words that the program cannot
    locate in the new text.
    """
    tags = [el.tag for el in obj.iter() if isinstance(el.tag, str)]
    if f"{{{WPD}}}inline" in tags:
        return False
    if f"{{{WPD}}}anchor" in tags:
        return True
    for el in obj.iter():
        if isinstance(el.tag, str) and el.tag.startswith(f"{{{VML}}}"):
            if "position:absolute" in (el.get("style") or "").replace(" ", "").lower():
                return True
    return False


def _nearest_paragraph(el):
    node = el.getparent()
    while node is not None and node.tag != _w("p"):
        node = node.getparent()
    return node


def _field_spans(p_el) -> list:
    """(begin, end) child positions of each outermost complex field in the paragraph.

    Complex fields (reference-manager citations, page references) are runs from a
    fldChar "begin" to a fldChar "end". Fields that do not close inside the
    paragraph are left out.
    """
    spans, stack = [], []
    for i, child in enumerate(p_el):
        if child.tag != _w("r"):
            continue
        for fc in child.findall("w:fldChar", NSMAP):
            kind = fc.get(_w("fldCharType"))
            if kind == "begin":
                stack.append(i)
            elif kind == "end" and stack:
                begin = stack.pop()
                if not stack:
                    spans.append((begin, i))
    return spans


SYMBOL_FONTS = frozenset(("symbol", "wingdings", "wingdings 2", "wingdings 3", "webdings", "mt extra",
                          "marlett", "zapfdingbats", "zapf dingbats"))

# Character formatting that changes what a reader sees. Where it differs between
# the runs a rewrite merges, merging them would change the visible text.
_SHOWN_PROPS = (("vanish", "hidden text"), ("strike", "struck-through text"),
                ("dstrike", "struck-through text"), ("caps", "capitals set by formatting"),
                ("smallCaps", "small capitals"), ("position", "raised or lowered text"),
                ("rStyle", "text in a different character style (for example, Emphasis)"))


def _prop_on(el) -> bool:
    if el is None:
        return False
    val = el.get(_w("val"))
    return val is None or val.lower() not in ("0", "false", "off")


def _shown_props(run_el) -> dict:
    rpr = run_el.find("w:rPr", NSMAP)
    sig = {}
    for tag, _ in _SHOWN_PROPS:
        el = rpr.find("w:" + tag, NSMAP) if rpr is not None else None
        if tag in ("position", "rStyle"):
            val = el.get(_w("val")) if el is not None else None
            sig[tag] = None if val in (None, "0", "") else val
        else:
            sig[tag] = _prop_on(el)
    return sig


def _in_symbol_font(run_el) -> bool:
    rpr = run_el.find("w:rPr", NSMAP)
    fonts = rpr.find("w:rFonts", NSMAP) if rpr is not None else None
    if fonts is None:
        return False
    return any((fonts.get(_w(a)) or "").strip().lower() in SYMBOL_FONTS for a in ("ascii", "hAnsi", "cs", "eastAsia"))


def _run_lock_reason(run_el):
    """Why a run makes its paragraph unsafe to rewrite, or None."""
    for child in run_el:
        if child.tag in OBJECT_TAGS:
            if not _object_floats(child):
                return "a picture or object placed between words"
        elif child.tag == _w("sym"):
            return "a symbol inserted from a symbol font"
        elif child.tag in (_w("tab"), _w("br"), _w("cr")):
            # the index shows the paragraph on one line, so the new text cannot carry these
            return "a tab or manual line break inside the text"
        elif child.tag not in _RUN_CHILDREN_OK:
            return "an element the program does not rewrite (%s)" % _local(child.tag)
    rpr = run_el.find("w:rPr", NSMAP)
    va = rpr.find("w:vertAlign", NSMAP) if rpr is not None else None
    if va is not None and va.get(_w("val")) in ("superscript", "subscript") and run_el.find("w:t", NSMAP) is not None:
        return "superscript or subscript text (for example, R squared)"
    return None


def paragraph_lock_reason(p_el):
    """Why a paragraph cannot be rewritten safely, or None when it can.

    Rewriting handles plain text runs (whose bold and italic it keeps), footnote
    and endnote marks, links, complex fields (kept intact while the text around
    them changes), and floating text boxes or pictures. A paragraph holding
    anything else is locked: tracked changes, content controls, simple fields,
    equations, symbol-font characters (inserted symbols or text in a symbol font),
    superscript or subscript text, pictures placed between words, fields that do
    not close inside the paragraph, any element the program does not know, and
    hidden, struck-through, capitalized (caps or small caps) or raised or lowered
    text, or a character style (such as Emphasis), set by formatting that differs
    between runs the rewrite would merge. A locked paragraph is left exactly as it is.
    """
    if count_tracked_changes(p_el):
        return "unresolved tracked changes"
    for child in p_el:
        tag = child.tag
        if not isinstance(tag, str):
            continue  # an XML comment or processing instruction holds no content
        if tag == _w("r"):
            reason = _run_lock_reason(child)
            if reason:
                return reason
        elif tag == _w("hyperlink"):
            for inner in child:
                if inner.tag == _w("r"):
                    if (_run_objects(inner) or inner.find("w:fldChar", NSMAP) is not None
                            or _run_has_footnote_ref(inner)):
                        return "a link that holds a field, a note mark or a picture"
                    reason = _run_lock_reason(inner)
                    if reason:
                        return reason
                elif inner.tag not in _NO_TEXT_CHILDREN:
                    return "a link that holds other elements"
        elif tag == _w("sdt"):
            return "a content control (for example, a citation inserted by a reference manager)"
        elif tag == _w("fldSimple"):
            return "a simple field (for example, a cross-reference)"
        elif isinstance(tag, str) and tag.startswith("{%s}" % MATH):
            return "an equation"
        elif tag not in _NO_TEXT_CHILDREN:
            return "an element the program does not rewrite (%s)" % _local(tag)
    depth = 0
    for run in p_el.findall("w:r", NSMAP):
        for fc in run.findall("w:fldChar", NSMAP):
            kind = fc.get(_w("fldCharType"))
            depth += 1 if kind == "begin" else (-1 if kind == "end" else 0)
            if depth < 0:
                return "a field that begins or ends outside this paragraph"
    if depth:
        return "a field that begins or ends outside this paragraph"
    for t_el in p_el.iter(_w("t")):
        if _nearest_paragraph(t_el) is not p_el:
            continue  # text inside a text box belongs to the box
        run = t_el.getparent()
        host = run.getparent() if run is not None else None
        if run is None or run.tag != _w("r") or host is None:
            return "text in an element the program does not rewrite"
        if host is p_el or (host.tag == _w("hyperlink") and host.getparent() is p_el):
            continue
        return "text in an element the program does not rewrite"
    # Formatting that changes what the text shows. The rewrite merges each stretch
    # of plain runs between links, fields and note marks into its first run, so a
    # property that differs inside a stretch would be lost; a symbol font would
    # turn a Greek letter into a Latin one wherever it sits.
    groups, _ = _tokenize_paragraph(p_el)
    for run in [r for g in groups for r in g] + [r for h in p_el.findall("w:hyperlink", NSMAP)
                                                 for r in h.findall("w:r", NSMAP)]:
        if _in_symbol_font(run):
            return "a character in a symbol font (for example, a Greek letter typed in Symbol)"
    for group in groups:
        sigs = [_shown_props(r) for r in group if get_run_text(r)]
        for tag, label in _SHOWN_PROPS:
            if len({s[tag] for s in sigs}) > 1:
                return label
    return None


NOTE_SEPARATOR_TYPES = ("separator", "continuationSeparator", "continuationNotice")
NOTE_NOT_FOUND = "[note text not found]"

def note_texts(doc) -> dict:
    """{(kind, id): text} for every footnote and endnote of a python-docx Document.

    kind is "footnote" or "endnote"; separators are left out, and a note of
    several paragraphs reads as one line, its paragraphs joined with a space.
    The texts are for reading: no program in this skill edits a note.
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    found = {}
    for rel in doc.part.rels.values():
        if rel.is_external or rel.reltype not in (RT.FOOTNOTES, RT.ENDNOTES):
            continue
        kind = "footnote" if rel.reltype == RT.FOOTNOTES else "endnote"
        root = etree.fromstring(rel.target_part.blob)
        for note in root.iter(_w(kind)):
            if note.get(_w("type")) in NOTE_SEPARATOR_TYPES:
                continue
            paras = ["".join(t.text or "" for t in p.iter(_w("t"))).strip() for p in note.iter(_w("p"))]
            found[(kind, note.get(_w("id")))] = " ".join(x for x in paras if x)
    return found


def paragraph_notes(p_el, texts) -> list:
    """[{kind, id, text}] for the footnote and endnote marks in a paragraph, in mark order.

    A mark whose note is missing from the file reads NOTE_NOT_FOUND, so that an
    empty note and a missing one are told apart.
    """
    notes = []
    for ref in p_el.iter(_w("footnoteReference"), _w("endnoteReference")):
        kind = "footnote" if ref.tag == _w("footnoteReference") else "endnote"
        nid = ref.get(_w("id"))
        notes.append({"kind": kind, "id": nid, "text": texts.get((kind, nid), NOTE_NOT_FOUND)})
    return notes


def _run_view_text(run_el) -> str:
    parts = []
    for child in run_el:
        tag = child.tag
        if tag == _w("t"):
            parts.append(child.text or "")
        elif tag == _w("tab"):
            parts.append("\t")
        elif tag in (_w("br"), _w("cr")):
            parts.append("\n")
        elif tag == _w("noBreakHyphen"):
            parts.append("-")
        elif tag in (_w("footnoteReference"), _w("endnoteReference")):
            parts.append("[^%s]" % child.get(_w("id"), "?"))
    return "".join(parts)


def paragraph_view_text(p_el) -> str:
    """The paragraph as the editors see it in the paragraph index.

    Runs and links in document order; tabs as tab characters; each footnote or
    endnote mark as [^N], N being the note's id. Text inside text boxes, content
    controls and simple fields is not part of the view (paragraphs holding
    content controls or simple fields are locked; see paragraph_lock_reason).
    """
    parts = []
    for child in p_el:
        if child.tag == _w("r"):
            parts.append(_run_view_text(child))
        elif child.tag == _w("hyperlink"):
            parts.extend(_run_view_text(r) for r in child.findall("w:r", NSMAP))
    return "".join(parts)



# ---------------------------------------------------------------------------
# Inline formatting restoration
# ---------------------------------------------------------------------------
# Single characters are normally too risky to restore: a lone "a" or "I" matches
# everywhere in the body text. The exception is statistical notation, where the
# italic single letter IS the convention (p < .05, n = 402, F(1, 385), M = 0.03).
# These are restored only when the original run was italic-not-bold, only on a
# word boundary, and case-sensitively, so "p" in "paper" is never touched.
SINGLE_CHAR_ITALIC_TOKENS = frozenset("npNrRtFdMzbBkQWU" "\u03b2\u03b3\u03c1\u03c3\u03c7\u03bc\u03c4\u03bb\u03b1\u03b4")


def _restore_single_char(body_text, clean, is_bold, is_italic, matches, used_ranges):
    """Boundary-checked, case-sensitive restoration for one statistical symbol."""
    start_from = 0
    while start_from < len(body_text):
        pos = body_text.find(clean, start_from)
        if pos < 0:
            break
        end = pos + len(clean)
        before_ok = pos == 0 or not (body_text[pos - 1].isalnum() or body_text[pos - 1] == "_")
        after_ok = end >= len(body_text) or not (body_text[end].isalnum() or body_text[end] == "_")
        overlap = any(pos < ue and end > us for us, ue in used_ranges)
        if before_ok and after_ok and not overlap:
            matches.append((pos, end, is_bold, is_italic))
            used_ranges.append((pos, end))
        start_from = end


def _restore_inline_formatting(body_run_el, formatted_catalog):
    """After basic replacement, find original formatted text in the body run
    and split the run to restore italic/bold formatting.

    formatted_catalog: list of (original_text, is_bold, is_italic) for runs
                       whose formatting was NOT preserved by prefix pull-back.
    """
    if not formatted_catalog:
        return

    body_text = get_run_text(body_run_el)
    if not body_text:
        return

    # Find matches: search for each formatted phrase in the body text
    matches = []  # (start_pos, end_pos, is_bold, is_italic)
    body_lower = body_text.lower()
    used_ranges = []

    # Sort by text length descending -- match longest phrases first
    sorted_catalog = sorted(formatted_catalog, key=lambda x: len(x[0]), reverse=True)

    for fmt_text, is_bold, is_italic in sorted_catalog:
        clean = fmt_text.strip()
        if len(clean) < 2:
            # Statistical symbols only; every other single character is skipped.
            if is_italic and not is_bold and clean in SINGLE_CHAR_ITALIC_TOKENS:
                _restore_single_char(body_text, clean, is_bold, is_italic, matches, used_ranges)
            continue

        search = clean.lower()
        # Try finding every occurrence (some phrases appear multiple times)
        start_from = 0
        while start_from < len(body_lower):
            pos = body_lower.find(search, start_from)
            if pos < 0:
                break
            end = pos + len(search)
            # Check overlap with existing matches
            overlap = any(pos < ue and end > us for us, ue in used_ranges)
            if not overlap:
                matches.append((pos, end, is_bold, is_italic))
                used_ranges.append((pos, end))
                start_from = end
            else:
                start_from = pos + 1

        # If exact match failed, try without trailing punctuation
        if not any(m for m in matches if body_lower[m[0]:m[1]] == search):
            search_stripped = search.rstrip('.,:; ')
            if len(search_stripped) >= 3 and search_stripped != search:
                start_from = 0
                while start_from < len(body_lower):
                    pos = body_lower.find(search_stripped, start_from)
                    if pos < 0:
                        break
                    end = pos + len(search_stripped)
                    overlap = any(pos < ue and end > us for us, ue in used_ranges)
                    if not overlap:
                        matches.append((pos, end, is_bold, is_italic))
                        used_ranges.append((pos, end))
                        start_from = end
                    else:
                        start_from = pos + 1

    if not matches:
        return

    matches.sort()

    # Build segments: [normal, formatted, normal, formatted, ...]
    segments = []  # (text, is_bold, is_italic)
    pos = 0
    for start, end, b, i in matches:
        if start > pos:
            segments.append((body_text[pos:start], False, False))
        segments.append((body_text[start:end], b, i))
        pos = end
    if pos < len(body_text):
        segments.append((body_text[pos:], False, False))

    # Skip if nothing meaningful to split
    if len(segments) <= 1:
        return

    # Skip empty segments
    segments = [(t, b, i) for t, b, i in segments if t]
    if not segments:
        return

    # First segment stays in the original run
    first_text, first_b, first_i = segments[0]
    _set_run_text(body_run_el, first_text)
    if first_b or first_i:
        _apply_formatting_to_run(body_run_el, first_b, first_i)

    # Create new runs for remaining segments
    para = body_run_el.getparent()
    current_idx = list(para).index(body_run_el)

    for seg_text, seg_bold, seg_italic in segments[1:]:
        # Properties and text only: a deep copy would repeat any tab, break or
        # picture the body run holds in every new segment.
        new_run = _plain_run_like(body_run_el, seg_text)
        _apply_formatting_to_run(new_run, seg_bold, seg_italic)
        current_idx += 1
        para.insert(current_idx, new_run)


# ---------------------------------------------------------------------------
# Paragraph text replacement (keeps prefix labels and restores bold and italic;
# see replace_paragraph_text for what else is kept and what locks a paragraph)
# ---------------------------------------------------------------------------
def _replace_segment_runs(text_runs: list, new_text: str) -> None:
    """Replace text across a list of runs, preserving formatting.

    Three-stage strategy:
    1. Catalog ALL formatted runs (text + bold/italic) BEFORE modification.
    2. Do prefix pull-back (formatted runs at the start of the paragraph).
    3. Restore inline formatting: find displaced formatted text in the body
       run and split to re-apply bold/italic.
    """
    if not text_runs:
        return

    if len(text_runs) == 1:
        _set_run_text(text_runs[0], new_text)
        return

    # STEP 0: Catalog ALL formatted runs BEFORE any modification
    full_catalog = []  # (text, is_bold, is_italic, run_index)
    for idx, run in enumerate(text_runs):
        if is_run_formatted(run):
            text = get_run_text(run)
            b, i = get_run_formatting(run)
            full_catalog.append((text, b, i, idx))

    # Classify runs
    normal_indices = [i for i, r in enumerate(text_runs) if not is_run_formatted(r)]
    formatted_indices = [i for i, r in enumerate(text_runs) if is_run_formatted(r)]

    # Track which catalog entries were preserved by prefix pull-back
    preserved_catalog_indices = set()

    if normal_indices:
        # --- CASE A: There are normal runs ---
        first_normal_idx = normal_indices[0]

        prefix_runs = [(i, text_runs[i]) for i in formatted_indices if i < first_normal_idx]
        body_runs = [text_runs[i] for i in range(first_normal_idx, len(text_runs))]

        if not prefix_runs:
            # No formatted prefix -- put in first run, clear rest
            _set_run_text(text_runs[0], new_text)
            for r in text_runs[1:]:
                _clear_run_text(r)
            # Restore inline formatting for ALL cataloged formatted runs
            to_restore = [(t, b, i) for t, b, i, _ in full_catalog]
            if to_restore:
                _restore_inline_formatting(text_runs[0], to_restore)
            return

        # Put ALL new text into the first body (normal) run, clear rest
        _set_run_text(body_runs[0], new_text)
        for r in body_runs[1:]:
            _clear_run_text(r)

        # Pull back prefix text from body run into formatted prefix runs
        remaining = new_text
        for run_idx, run_el in prefix_runs:
            orig_prefix_text = get_run_text(run_el)
            if remaining.startswith(orig_prefix_text):
                # Match! Keep this prefix run, trim from body
                remaining = remaining[len(orig_prefix_text):]
                # Mark as preserved in catalog
                for ci, (ct, cb, cit, cidx) in enumerate(full_catalog):
                    if cidx == run_idx:
                        preserved_catalog_indices.add(ci)
                        break
            else:
                # Prefix doesn't match (text was reorganized by editing)
                _clear_run_text(run_el)

        # Update body run with remaining text
        _set_run_text(body_runs[0], remaining)

        # Restore inline formatting for unpreserved formatted runs
        to_restore = [
            (t, b, i) for ci, (t, b, i, _) in enumerate(full_catalog)
            if ci not in preserved_catalog_indices
        ]
        if to_restore:
            _restore_inline_formatting(body_runs[0], to_restore)

    else:
        # --- CASE B: ALL runs are formatted (no normal runs) ---
        fmts = [get_run_formatting(r) for r in text_runs]

        if len(set(fmts)) == 1:
            # All same format (e.g., all-italic hypothesis)
            _set_run_text(text_runs[0], new_text)
            for r in text_runs[1:]:
                _clear_run_text(r)
            return

        # Mixed formats -- find where format changes (prefix boundary)
        first_fmt = fmts[0]
        prefix_end = 0
        for i, fmt in enumerate(fmts):
            if fmt != first_fmt:
                prefix_end = i
                break

        if prefix_end == 0:
            _set_run_text(text_runs[0], new_text)
            for r in text_runs[1:]:
                _clear_run_text(r)
            return

        prefix_runs_list = text_runs[:prefix_end]
        body_runs_list = text_runs[prefix_end:]
        prefix_text = "".join(get_run_text(r) for r in prefix_runs_list)

        # Catalog inline formatted from body that differ from body's primary format
        body_primary_fmt = fmts[prefix_end]
        inline_formatted = []
        for r in body_runs_list:
            fmt = get_run_formatting(r)
            if fmt != body_primary_fmt:
                t = get_run_text(r)
                inline_formatted.append((t, fmt[0], fmt[1]))

        if new_text.startswith(prefix_text):
            body_text = new_text[len(prefix_text):]
            _set_run_text(body_runs_list[0], body_text)
            for r in body_runs_list[1:]:
                _clear_run_text(r)
        else:
            # Prefix doesn't match -- catalog prefix for inline restoration
            for r in prefix_runs_list:
                t = get_run_text(r)
                b, i = get_run_formatting(r)
                inline_formatted.append((t, b, i))
            _set_run_text(body_runs_list[0], new_text)
            for r in body_runs_list[1:]:
                _clear_run_text(r)
            for r in prefix_runs_list:
                _clear_run_text(r)

        if inline_formatted:
            _restore_inline_formatting(body_runs_list[0], inline_formatted)


def _field_result_text(field_children) -> str:
    """Displayed text of a complex field: what shows between its "separate" and "end" marks."""
    parts, stack = [], []

    def visit(run):
        fc = run.find("w:fldChar", NSMAP)
        if fc is not None:
            kind = fc.get(_w("fldCharType"))
            if kind == "begin":
                stack.append("instruction")
            elif kind == "separate" and stack:
                stack[-1] = "result"
            elif kind == "end" and stack:
                stack.pop()
            return
        if stack and all(state == "result" for state in stack):
            parts.append(_run_view_text(run))

    for el in field_children:
        if el.tag == _w("r"):
            visit(el)
        else:
            for run in el.iter(_w("r")):
                visit(run)
    return "".join(parts)


def _tokenize_paragraph(p_el):
    """Split a paragraph into groups of plain text runs and the fixed tokens between them.

    Returns (groups, tokens) with len(groups) == len(tokens) + 1: groups[i] holds
    the text runs before tokens[i]. A token is ("note", id) for a footnote or
    endnote mark (its [^id] must appear at that point in the new text), or
    ("literal", text) for a link or a complex field whose displayed words must
    appear unchanged in the new text. Runs without text (field marks, comment
    references) and floating objects stay where they are and belong to no group.
    """
    children = list(p_el)
    in_field = {}
    for begin, end in _field_spans(p_el):
        for i in range(begin, end + 1):
            in_field[i] = (begin, end)
    groups, tokens = [[]], []
    i = 0
    while i < len(children):
        child = children[i]
        if i in in_field:
            begin, end = in_field[i]
            shown = _field_result_text(children[begin:end + 1])
            if shown:
                tokens.append(("literal", shown))
                groups.append([])
            i = end + 1
            continue
        if child.tag == _w("hyperlink"):
            shown = "".join(_run_view_text(r) for r in child.findall("w:r", NSMAP))
            if shown:
                tokens.append(("literal", shown))
                groups.append([])
        elif child.tag == _w("r"):
            if _run_has_footnote_ref(child):
                ref = child.find("w:footnoteReference", NSMAP)
                if ref is None:
                    ref = child.find("w:endnoteReference", NSMAP)
                tokens.append(("note", ref.get(_w("id"), "?")))
                groups.append([])
            elif child.findall("w:t", NSMAP):
                groups[-1].append(child)
        i += 1
    return groups, tokens


def _split_new_text(new_text: str, tokens: list) -> list:
    """Cut the new text around the paragraph's fixed tokens; ValueError if it cannot be.

    Each note mark must come back with its own number, in the paragraph's order.
    Link and field words are matched with curly and straight quotes treated alike.
    """
    pieces, pos = [], 0
    folded = _fold_quotes(new_text)
    for kind, shown in tokens:
        if kind == "note":
            m = _NOTE_MARK.search(new_text, pos)
            if m is None:
                raise ValueError("paragraph not rewritten: a footnote or endnote mark ([^N]) "
                                 "is missing from the new text")
            if m.group(0) != "[^%s]" % shown:
                raise ValueError("paragraph not rewritten: footnote or endnote marks were moved, "
                                 "reordered or renumbered (expected [^%s], found %s)" % (shown, m.group(0)))
            pieces.append(new_text[pos:m.start()])
            pos = m.end()
        else:
            k = folded.find(_fold_quotes(shown), pos)
            if k < 0:
                raise ValueError("paragraph not rewritten: the new text changes or drops the words "
                                 "of a link or citation field (%s)" % shown[:60])
            pieces.append(new_text[pos:k])
            pos = k + len(shown)
    pieces.append(new_text[pos:])
    if any(_NOTE_MARK.search(piece) for piece in pieces):
        raise ValueError("paragraph not rewritten: the new text has more footnote or endnote "
                         "marks than the paragraph")
    return pieces


def _restore_children(p_el, saved) -> None:
    """Put a paragraph back as it was (saved is a deep copy taken before the rewrite)."""
    for child in list(p_el):
        p_el.remove(child)
    for child in list(saved):
        p_el.append(child)
    for key in list(p_el.attrib):
        del p_el.attrib[key]
    for key, value in saved.attrib.items():
        p_el.set(key, value)


def _comment_words_before(p_el) -> dict:
    """{comment id: the words it covers} for comments whose start and end marks both
    sit directly in the paragraph, read before a rewrite."""
    covered, done, open_ids = {}, {}, []
    for child in p_el:
        if child.tag == _w("commentRangeStart"):
            cid = child.get(_w("id"))
            open_ids.append(cid)
            covered[cid] = ""
        elif child.tag == _w("commentRangeEnd"):
            cid = child.get(_w("id"))
            if cid in open_ids:
                open_ids.remove(cid)
                done[cid] = covered[cid]
        elif open_ids:
            text = "".join(t.text or "" for t in child.iter(_w("t")) if _nearest_paragraph(t) is p_el)
            for cid in open_ids:
                covered[cid] += text
    return done


def _reattach_existing_comments(p_el, covered: dict) -> set:
    """Put each earlier comment back on the sentence holding the words it covered,
    when those words survive the rewrite unchanged and appear once. Returns the ids
    re-attached; the others are widened by _widen_existing_comments."""
    done = set()
    for cid, words in covered.items():
        start = next((el for el in p_el if el.tag == _w("commentRangeStart") and el.get(_w("id")) == cid), None)
        end = next((el for el in p_el if el.tag == _w("commentRangeEnd") and el.get(_w("id")) == cid), None)
        if start is None or end is None or not words.strip():
            continue
        own, _ = _paragraph_char_map(p_el)
        found, _count = _find_phrase(own, words)
        if found is None:
            continue
        ref = end.getnext()
        has_ref = (ref is not None and ref.tag == _w("r")
                   and any(c.get(_w("id")) == cid for c in ref.findall("w:commentReference", NSMAP)))
        p_el.remove(start)
        p_el.remove(end)
        if has_ref:
            p_el.remove(ref)
        covering = [(s, e) for s, e in sentence_spans(own) if s < found[1] and e > found[0]]
        span = (covering[0][0], covering[-1][1]) if covering else found
        _, entries = _paragraph_char_map(p_el)
        t_end, k_end = _entry_at(entries, span[1], "end")
        _end_anchor(p_el, t_end, k_end).addnext(end)
        if has_ref:
            end.addnext(ref)
        _, entries = _paragraph_char_map(p_el)
        t_start, k_start = _entry_at(entries, span[0], "start")
        _start_anchor(p_el, t_start, k_start).addprevious(start)
        done.add(cid)
    return done


def _widen_existing_comments(p_el, skip=frozenset()) -> list:
    """Move the range marks of comments already in a rewritten paragraph to its two ends.

    After a rewrite the words sit in different runs, so a mark left in place would
    tie an author's earlier comment to other words. Widened, the comment covers
    the whole paragraph. Comments in `skip` (already re-attached to their
    sentence) are left alone. Returns a note when any mark moved.
    """
    def is_text(el):
        return el.tag == _w("hyperlink") or (el.tag == _w("r") and el.find("w:t", NSMAP) is not None)

    moved = False
    for start in [el for el in p_el if el.tag == _w("commentRangeStart") and el.get(_w("id")) not in skip]:
        kids = list(p_el)
        if any(is_text(el) for el in kids[:kids.index(start)]):
            p_el.remove(start)
            first = next((el for el in p_el if el.tag != _w("pPr")), None)
            if first is not None:
                first.addprevious(start)
            else:
                p_el.append(start)
            moved = True
    for end in [el for el in p_el if el.tag == _w("commentRangeEnd") and el.get(_w("id")) not in skip]:
        kids = list(p_el)
        if any(is_text(el) for el in kids[kids.index(end) + 1:]):
            cid = end.get(_w("id"))
            ref = end.getnext()
            has_ref = (ref is not None and ref.tag == _w("r")
                       and any(c.get(_w("id")) == cid for c in ref.findall("w:commentReference", NSMAP)))
            p_el.remove(end)
            p_el.append(end)
            if has_ref:
                p_el.remove(ref)
                p_el.append(ref)
            moved = True
    return ["an existing comment in this paragraph now covers the whole paragraph"] if moved else []


def replace_paragraph_text(para, new_text: str) -> list:
    """Replace the visible text of a paragraph; all or nothing.

    - Refuses a locked paragraph (paragraph_lock_reason: tracked changes, content
      controls, simple fields, equations, symbol-font characters, superscript or
      subscript text, pictures placed between words, ...): ValueError, and the
      paragraph is untouched.
    - Converts straight quotes to curly quotes.
    - Keeps each footnote or endnote mark after the words it follows: the new text
      carries each [^N] the editors saw (paragraph_view_text), same numbers, same order.
    - Keeps links and reference-manager (complex) fields intact: the new text must
      contain their displayed words unchanged and in the same order.
    - Leaves floating text boxes and pictures alone.
    - Keeps bold and italic for prefix labels ("Participants.", "Measures.") and
      restores them for inline phrases and statistical symbols. Other differences
      in character formatting between the rewritten runs (font, size, color,
      underline, highlight) are not kept: the new words take the first run's.
    - Keeps a comment that already sits in the paragraph on the sentence holding
      the words it covered, when those words survive unchanged; otherwise widens
      it to the whole paragraph. Either way the returned notes say so.
    - Checks the result: if the paragraph would not read exactly as the new text,
      it is put back as it was and ValueError is raised.

    Returns a list of notes for the record (empty when there is nothing to note).
    """
    p_el = para._element
    reason = paragraph_lock_reason(p_el)
    if reason:
        raise ValueError("paragraph not rewritten: it holds %s" % reason)

    new_text = smartify_quotes(new_text)
    groups, tokens = _tokenize_paragraph(p_el)
    pieces = _split_new_text(new_text, tokens)  # raises before anything changes
    for group, piece in zip(groups, pieces):
        if piece and not group:
            raise ValueError("paragraph not rewritten: the new text puts words where the "
                             "paragraph has no text run to hold them")

    saved = copy.deepcopy(p_el)
    covered = _comment_words_before(p_el)
    try:
        for group, piece in zip(groups, pieces):
            if group:
                _replace_segment_runs(group, piece)
        after = paragraph_view_text(p_el)
        if _fold_quotes(after) != _fold_quotes(new_text):
            raise ValueError("paragraph not rewritten: the result would not read as the new text "
                             "(for example, a tab or line break inside a run)")
        kept = _reattach_existing_comments(p_el, covered)
        notes = ["an existing comment in this paragraph stays on its sentence"] * len(kept)
        return notes + _widen_existing_comments(p_el, skip=kept)
    except Exception:
        _restore_children(p_el, saved)
        raise


# ---------------------------------------------------------------------------
# Sentences
# ---------------------------------------------------------------------------
_SENTENCE_END = re.compile("[.!?]+[\"'\u201d\u2019)\\]]*")
_CLOSERS = "\"'\u201d\u2019)]"
_OPENERS = "\"'\u201c\u2018(["
_ABBREVIATIONS = frozenset((
    "e.g", "i.e", "al", "cf", "vs", "viz", "resp", "approx", "ca",
    "fig", "figs", "eq", "eqs", "no", "nos", "vol", "vols", "p", "pp", "ch", "sec",
    "dr", "mr", "ms", "mrs", "prof", "st", "mt", "jr", "sr",
    "inc", "corp", "ltd", "co", "u.s", "u.k", "ph.d"))


def _period_is_not_an_end(text: str, dot: int) -> bool:
    """True when the period at text[dot] belongs to an abbreviation or an initial."""
    before = text[:dot]
    m = re.search(r"([A-Za-z][A-Za-z.]*)$", before)
    if not m:
        return False  # after a digit, a bracket or a quote: a real end
    token = m.group(1)
    if token.lower().rstrip(".") in _ABBREVIATIONS:
        return True
    if len(token) == 1 and token.isupper():
        # A single capital letter is an initial in a name ("J. Smith"), unless a
        # lowercase word precedes it ("found X. Y follows").
        words = before[:m.start()].split()
        if not words:
            return True
        prev = words[-1]
        if prev.endswith(".") and len(prev.rstrip(".")) == 1:
            return True  # another initial: "M. J. Quill"
        return prev[0].isupper() or not prev[0].isalpha()
    return False


def sentence_spans(text: str) -> list:
    """(start, end) character spans of the sentences in a paragraph's text.

    A sentence ends at ".", "?" or "!" (with any closing quotes or brackets) that
    is followed by white space and then a capital letter, a digit, an opening
    quote or bracket. A period after a common abbreviation ("e.g.", "et al.",
    "Fig.", "U.S."), inside a number ("0.42", "p < .05"), or after an initial in
    a name ("J. Smith") does not end a sentence. Leading and trailing white space
    is outside every span. Known limit: a single capital letter after a
    capitalized word ("Appendix A. The") reads as an initial, so those two
    sentences stay together.
    """
    n = len(text)
    start = 0
    while start < n and text[start].isspace():
        start += 1
    spans = []
    for m in _SENTENCE_END.finditer(text):
        end = m.end()
        if end >= n or not text[end].isspace():
            continue
        k = end
        while k < n and text[k].isspace():
            k += 1
        if k >= n:
            continue
        nxt = text[k]
        if not (nxt.isupper() or nxt.isdigit() or nxt in _OPENERS):
            continue
        punct = m.group(0)
        if punct[0] == "." and len(punct.rstrip(_CLOSERS)) == 1 and _period_is_not_an_end(text, m.start()):
            continue
        if end > start:
            spans.append((start, end))
        start = k
    stop = n
    while stop > start and text[stop - 1].isspace():
        stop -= 1
    if stop > start:
        spans.append((start, stop))
    return spans


# ---------------------------------------------------------------------------
# Finding a phrase and placing a comment range around its sentence
# ---------------------------------------------------------------------------
_MATCH_MAP = {0x2018: "'", 0x2019: "'", 0x201A: "'", 0x201B: "'", 0x201C: '"', 0x201D: '"',
              0x201E: '"', 0x201F: '"', 0x2013: "-", 0x2014: "-", 0x2012: "-", 0x2010: "-",
              0x2011: "-", 0x00A0: " ", 0x2009: " ", 0x202F: " "}


def _normalize_with_positions(s: str):
    """Text with curly quotes, dashes and runs of white space normalized, plus the
    original position of every character kept, so a match maps back exactly."""
    out, positions, prev_space = [], [], False
    for i, ch in enumerate(s):
        c = _MATCH_MAP.get(ord(ch), ch)
        if c.isspace():
            if prev_space:
                continue
            c, prev_space = " ", True
        else:
            prev_space = False
        out.append(c)
        positions.append(i)
    return "".join(out), positions


def _find_phrase(haystack: str, phrase: str):
    """((start, end), count): where a phrase sits in the text, matched across curly and
    straight quotes, dashes and white space; (None, count) when absent or repeated."""
    needle, _ = _normalize_with_positions(_NOTE_MARK.sub("", phrase or "").strip())
    needle = needle.strip()
    if not needle:
        return None, 0
    hay, positions = _normalize_with_positions(haystack)

    def hits_of(n):
        found, k = [], hay.find(n)
        while k >= 0:
            found.append(k)
            k = hay.find(n, k + 1)
        return found

    hits = hits_of(needle)
    if not hits and len(needle) > 2 and needle[0] in "\"'" and needle[-1] in "\"'":
        needle = needle[1:-1].strip()  # a phrase copied with its quotation marks around it
        hits = hits_of(needle) if needle else []
    if len(hits) != 1:
        return None, len(hits)
    k = hits[0]
    return (positions[k], positions[k + len(needle) - 1] + 1), 1


def phrase_count(text: str, phrase: str) -> int:
    """How many times a phrase occurs in a text, counted the way a comment is
    placed: curly and straight quotes, dashes and white space normalized, and
    overlapping occurrences counted. A comment lands on its sentence only when
    this is 1."""
    return _find_phrase(text, phrase)[1]


def _paragraph_char_map(p_el):
    """The paragraph's own text (every <w:t> whose nearest paragraph is this one,
    so text-box text is left out) and each <w:t> with its (start, end) offsets."""
    entries, parts, pos = [], [], 0
    for t_el in p_el.iter(_w("t")):
        if _nearest_paragraph(t_el) is not p_el:
            continue
        s = t_el.text or ""
        entries.append((t_el, pos, pos + len(s)))
        parts.append(s)
        pos += len(s)
    return "".join(parts), entries


def _entry_at(entries, pos: int, side: str):
    """The <w:t> holding the first character at pos (side "start") or the last
    character before pos (side "end"), with the offset inside it."""
    for t_el, a, b in entries:
        if (side == "start" and a <= pos < b) or (side == "end" and a < pos <= b):
            return t_el, pos - a
    raise ValueError("offset %d is outside the paragraph text" % pos)


def _split_run(run, t_el, k: int):
    """Move t_el's text from offset k on, and every later child of the run, into a new
    run with the same properties, placed right after the run. Returns the new run."""
    new_run = etree.Element(_w("r"))
    rpr = run.find("w:rPr", NSMAP)
    if rpr is not None:
        new_run.append(copy.deepcopy(rpr))
    children = list(run)
    idx = children.index(t_el)
    text = t_el.text or ""
    if k <= 0:
        movers = children[idx:]
    else:
        movers = children[idx + 1:]
        if k < len(text):
            tail = etree.SubElement(new_run, _w("t"))
            tail.text = text[k:]
            tail.set(XML_SPACE, "preserve")
            t_el.text = text[:k]
            t_el.set(XML_SPACE, "preserve")
    for child in movers:
        new_run.append(child)
    run.addnext(new_run)
    return new_run


def _child_of_paragraph(p_el, el):
    while el.getparent() is not p_el:
        el = el.getparent()
    return el


def _span_holding(spans, idx: int):
    for begin, end in spans:
        if begin <= idx <= end:
            return begin, end
    return None


def _start_anchor(p_el, t_el, k: int):
    """The paragraph child before which a range starting at offset k of t_el begins.

    A start inside a link, content control or simple field moves to before that
    element; a start inside a complex field moves to before its "begin" run, so a
    reference-manager refresh cannot remove the marker.
    """
    run = t_el.getparent()
    spans = _field_spans(p_el)
    if run.getparent() is p_el:
        idx = list(p_el).index(run)
        span = _span_holding(spans, idx)
        if span and span[0] < idx:
            return list(p_el)[span[0]]
        content = [c for c in run if c.tag != _w("rPr")]
        if k == 0 and content and content[0] is t_el:
            return run
        return _split_run(run, t_el, k)
    el = _child_of_paragraph(p_el, run)
    span = _span_holding(spans, list(p_el).index(el))
    return list(p_el)[span[0]] if span else el


def _end_anchor(p_el, t_el, k: int):
    """The paragraph child after which a range ending at offset k of t_el ends
    (the mirror of _start_anchor)."""
    run = t_el.getparent()
    spans = _field_spans(p_el)
    if run.getparent() is p_el:
        idx = list(p_el).index(run)
        span = _span_holding(spans, idx)
        if span and idx < span[1]:
            return list(p_el)[span[1]]
        content = [c for c in run if c.tag != _w("rPr")]
        if k >= len(t_el.text or "") and content and content[-1] is t_el:
            return run
        _split_run(run, t_el, k)
        return run
    el = _child_of_paragraph(p_el, run)
    span = _span_holding(spans, list(p_el).index(el))
    return list(p_el)[span[1]] if span else el


def _comment_reference_run(cid: str):
    ref_r = etree.Element(_w("r"))
    ref_rpr = etree.SubElement(ref_r, _w("rPr"))
    ref_style = etree.SubElement(ref_rpr, _w("rStyle"))
    ref_style.set(_w("val"), "CommentReference")
    ref = etree.SubElement(ref_r, _w("commentReference"))
    ref.set(_w("id"), cid)
    return ref_r


def _mark_comment_range(p_el, cid: str, highlight):
    """Insert the range markers and the reference run for one comment.

    With a highlight phrase found once in the paragraph's own text, the range
    covers the sentence (or sentences) holding it; otherwise the whole paragraph.
    Returns a note when the comment fell back to the whole paragraph, else None.
    Never changes the paragraph's text.
    """
    start_el = etree.Element(_w("commentRangeStart"))
    start_el.set(_w("id"), cid)
    end_el = etree.Element(_w("commentRangeEnd"))
    end_el.set(_w("id"), cid)
    ref_run = _comment_reference_run(cid)

    note, span = None, None
    if highlight:
        text, _ = _paragraph_char_map(p_el)
        found, count = _find_phrase(text, highlight)
        if found is None:
            note = ("highlight phrase not found in the paragraph" if count == 0
                    else "highlight phrase occurs %d times in the paragraph" % count)
            note += "; the comment covers the whole paragraph"
        else:
            covering = [(s, e) for s, e in sentence_spans(text) if s < found[1] and e > found[0]]
            span = (covering[0][0], covering[-1][1]) if covering else found

    if span:
        # The end first: splitting a run at the end leaves every earlier offset in place.
        _, entries = _paragraph_char_map(p_el)
        t_end, k_end = _entry_at(entries, span[1], "end")
        after = _end_anchor(p_el, t_end, k_end)
        after.addnext(end_el)
        end_el.addnext(ref_run)
        _, entries = _paragraph_char_map(p_el)
        t_start, k_start = _entry_at(entries, span[0], "start")
        _start_anchor(p_el, t_start, k_start).addprevious(start_el)
        return note

    content = [c for c in p_el if c.tag != _w("pPr")]
    if content:
        content[0].addprevious(start_el)
    else:
        p_el.append(start_el)
    p_el.append(end_el)
    p_el.append(ref_run)
    return note


# ---------------------------------------------------------------------------
# Word comment injection via XML
# ---------------------------------------------------------------------------
def inject_comments_ex(docx_path: Path, items: list, date: str = None):
    """Add Word comments to a .docx (modified in place).

    Each item is a dict:
        text       the comment text
        author     the comment author
        search     a phrase the target paragraph contains (first paragraph that
                   contains it, text boxes and table cells included), or
        para_index the target's position among the body's top-level paragraphs
                   (python-docx's Document.paragraphs order)
        highlight  optional: a phrase in the target paragraph; the comment then
                   covers the sentence holding it, else the whole paragraph

    Existing comments are kept and new ids continue from the highest one. Range
    markers never sit inside a link, a content control, a simple field, or a
    complex field that opens and closes in the paragraph (a field spanning
    several paragraphs, such as a reference-manager bibliography, is not
    detected).

    Returns (added, notes): the number of comments added, and a list of
    (item index, message) for each comment not added or placed on the whole
    paragraph instead of a sentence.
    """
    if date is None:
        date = "%sT00:00:00Z" % datetime.date.today().isoformat()
    docx_path = Path(docx_path)
    buf = BytesIO(docx_path.read_bytes())
    zin = zipfile.ZipFile(buf, "r")
    doc_xml = etree.fromstring(zin.read("word/document.xml"))

    # Preserve existing comments: read comments.xml if it exists
    existing_comments_xml = None
    if "word/comments.xml" in zin.namelist():
        existing_comments_xml = zin.read("word/comments.xml")

    if existing_comments_xml is not None:
        comments_root = etree.fromstring(existing_comments_xml)
        # Find max existing comment ID to continue numbering
        max_existing_id = -1
        for c_elem in comments_root.findall(f"{{{WML}}}comment"):
            cid_str = c_elem.get(f"{{{WML}}}id", "0")
            try:
                max_existing_id = max(max_existing_id, int(cid_str))
            except ValueError:
                pass
        comment_id = max_existing_id + 1
    else:
        comments_root = etree.Element(
            f"{{{WML}}}comments",
            nsmap={"w": WML, "r": REL},
        )
        comment_id = 0

    added, notes = 0, []
    body = doc_xml.find("w:body", NSMAP)
    top_paragraphs = body.findall("w:p", NSMAP) if body is not None else []

    for n, item in enumerate(items):
        target_p = None
        if item.get("para_index") is not None:
            pos = item["para_index"]
            if 0 <= pos < len(top_paragraphs):
                target_p = top_paragraphs[pos]
        else:
            search_text = item.get("search") or ""
            for p_elem in doc_xml.iter(_w("p")):
                p_text = "".join(t.text or "" for t in p_elem.iter(_w("t")))
                if search_text and search_text in p_text:
                    target_p = p_elem
                    break

        if target_p is None:
            notes.append((n, "paragraph not found; comment not added"))
            continue

        cid = str(comment_id)

        c_elem = etree.SubElement(comments_root, _w("comment"))
        c_elem.set(_w("id"), cid)
        c_elem.set(_w("author"), item.get("author") or "")
        c_elem.set(_w("date"), date)
        c_p = etree.SubElement(c_elem, _w("p"))
        c_r = etree.SubElement(c_p, _w("r"))
        c_t = etree.SubElement(c_r, _w("t"))
        c_t.text = item.get("text") or ""
        c_t.set(XML_SPACE, "preserve")

        note = _mark_comment_range(target_p, cid, item.get("highlight"))
        if note:
            notes.append((n, note))

        comment_id += 1
        added += 1

    if added == 0:
        zin.close()
        return 0, notes

    out_buf = BytesIO()
    zout = zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED)

    for item in zin.infolist():
        if item.filename == "word/document.xml":
            zout.writestr(
                item,
                etree.tostring(doc_xml, xml_declaration=True, encoding="UTF-8", standalone=True),
            )
        elif item.filename == "[Content_Types].xml":
            ct_xml = etree.fromstring(zin.read(item.filename))
            ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
            existing = [e for e in ct_xml if e.get("PartName") == "/word/comments.xml"]
            if not existing:
                override = etree.SubElement(ct_xml, f"{{{ct_ns}}}Override")
                override.set("PartName", "/word/comments.xml")
                override.set(
                    "ContentType",
                    "application/vnd.openxmlformats-officedocument"
                    ".wordprocessingml.comments+xml",
                )
            zout.writestr(
                item,
                etree.tostring(ct_xml, xml_declaration=True, encoding="UTF-8", standalone=True),
            )
        elif item.filename == "word/_rels/document.xml.rels":
            rels_xml = etree.fromstring(zin.read(item.filename))
            rels_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
            existing_rels = [e for e in rels_xml if "comments" in (e.get("Type") or "")]
            if not existing_rels:
                max_id = 0
                for e in rels_xml:
                    rid = e.get("Id", "")
                    if rid.startswith("rId"):
                        try:
                            max_id = max(max_id, int(rid[3:]))
                        except ValueError:
                            pass
                new_rel = etree.SubElement(rels_xml, f"{{{rels_ns}}}Relationship")
                new_rel.set("Id", f"rId{max_id + 1}")
                new_rel.set(
                    "Type",
                    "http://schemas.openxmlformats.org/officeDocument"
                    "/2006/relationships/comments",
                )
                new_rel.set("Target", "comments.xml")
            zout.writestr(
                item,
                etree.tostring(rels_xml, xml_declaration=True, encoding="UTF-8", standalone=True),
            )
        elif item.filename == "word/comments.xml":
            # Skip -- we write the merged comments_root below
            pass
        else:
            zout.writestr(item, zin.read(item.filename))

    zout.writestr(
        "word/comments.xml",
        etree.tostring(comments_root, xml_declaration=True, encoding="UTF-8", standalone=True),
    )

    zout.close()
    zin.close()
    docx_path.write_bytes(out_buf.getvalue())
    return added, notes


def inject_comments(docx_path: Path, comments: list, date: str = None) -> int:
    """Add Word comments to a .docx (modified in place); returns the number added.

    comments: (search_text, comment_text, author) tuples, or
              (search_text, comment_text, author, highlight) to cover the sentence
              holding the highlight phrase instead of the whole paragraph.
    See inject_comments_ex for placement rules and the notes it can return.
    """
    items = []
    for c in comments:
        items.append({"search": c[0], "text": c[1], "author": c[2],
                      "highlight": c[3] if len(c) > 3 else None})
    added, _ = inject_comments_ex(docx_path, items, date=date)
    return added


# ---------------------------------------------------------------------------
# Verification
# ---------------------------------------------------------------------------
def verify_docx_formatting(docx_path: Path, original_path: Path = None) -> dict:
    """Verify a .docx output for common formatting errors.

    Checks:
    1. Straight quotes, double and single (none added relative to the original;
       zero if no original is given)
    2. Inline formatting loss (original had mixed formatting, output is uniform)
    3. Footnote preservation (if original provided)
    4. Table preservation (if original provided)
    """
    from docx import Document

    errors = []
    warnings = []

    doc = Document(str(docx_path))

    # --- Check 1: Straight quotes (double and single, including apostrophes) ---
    # The totals are compared with the original: a net increase is an error (the
    # edits introduced straight quotes); straight quotes the original already
    # had are reported as a warning.
    def count_straight(d):
        return sum(p.text.count('"') + p.text.count("'") for p in d.paragraphs)

    docx_path = Path(docx_path)
    original_path = Path(original_path) if original_path else None
    straight = count_straight(doc)
    orig_straight = None
    if original_path and original_path.exists():
        orig_straight = count_straight(Document(str(original_path)))

    if orig_straight is None:
        if straight > 0:
            errors.append(f"Found {straight} straight quotes or apostrophes. All should be curly.")
    elif straight > orig_straight:
        errors.append(f"Found {straight} straight quotes or apostrophes; the original had {orig_straight}. "
                      f"Edited text should use curly quotes.")
    elif straight > 0:
        warnings.append(f"{straight} straight quotes or apostrophes remain from the original "
                        f"(which had {orig_straight}); none were introduced by editing.")

    # --- Check 2: Compare with original ---
    if original_path and original_path.exists():
        orig_doc = Document(str(original_path))

        # Footnote count
        orig_xml = etree.tostring(orig_doc.element)
        out_xml = etree.tostring(doc.element)
        orig_fn = orig_xml.count(b"footnoteReference")
        out_fn = out_xml.count(b"footnoteReference")
        if orig_fn != out_fn:
            errors.append(f"Footnote count mismatch: original={orig_fn}, output={out_fn}")

        # Table count
        if len(orig_doc.tables) != len(doc.tables):
            errors.append(
                f"Table count mismatch: original={len(orig_doc.tables)}, "
                f"output={len(doc.tables)}"
            )

        # Check for inline formatting loss
        def get_unique_fmts(para):
            fmts = set()
            for r in para._element.findall("./w:r", NSMAP):
                t = get_run_text(r)
                if t.strip():
                    b, it = get_run_formatting(r)
                    if b and it:
                        fmts.add("BI")
                    elif b:
                        fmts.add("B")
                    elif it:
                        fmts.add("I")
                    else:
                        fmts.add("N")
            return fmts

        orig_count = len(orig_doc.paragraphs)
        out_count = len(doc.paragraphs)
        counts_match = (orig_count == out_count)

        if not counts_match:
            warnings.append(
                f"Paragraph count changed: original={orig_count}, output={out_count}. "
                f"Using text-based matching for formatting verification."
            )

        fmt_loss_count = 0

        if counts_match:
            # Index-based comparison (fast, exact)
            for pi in range(orig_count):
                orig_fmts = get_unique_fmts(orig_doc.paragraphs[pi])
                out_fmts = get_unique_fmts(doc.paragraphs[pi])
                if len(orig_fmts) > 1 and len(out_fmts) == 1:
                    fmt_loss_count += 1
                    if fmt_loss_count <= 5:
                        errors.append(
                            f"P{pi:03d}: Original had formats {orig_fmts}, output has only {out_fmts}. "
                            f"Inline formatting lost."
                        )
        else:
            # Text-based matching: build index of original paragraphs by text prefix
            orig_index = {}  # text_prefix -> (para_index, para_obj)
            for pi, para in enumerate(orig_doc.paragraphs):
                text = para.text.strip()
                if len(text) >= 20:
                    key = text[:80]
                    if key not in orig_index:
                        orig_index[key] = (pi, para)

            for oi, out_para in enumerate(doc.paragraphs):
                out_text = out_para.text.strip()
                if len(out_text) < 20:
                    continue
                out_key = out_text[:80]
                matched = orig_index.get(out_key)
                if matched is None:
                    continue
                orig_pi, orig_para = matched
                orig_fmts = get_unique_fmts(orig_para)
                out_fmts = get_unique_fmts(out_para)
                if len(orig_fmts) > 1 and len(out_fmts) == 1:
                    fmt_loss_count += 1
                    if fmt_loss_count <= 5:
                        errors.append(
                            f"P{oi:03d} (matched original P{orig_pi:03d}): Original had formats "
                            f"{orig_fmts}, output has only {out_fmts}. Inline formatting lost."
                        )

        if fmt_loss_count > 5:
            errors.append(f"... and {fmt_loss_count - 5} more inline formatting losses.")

    passed = len(errors) == 0
    return {"passed": passed, "errors": errors, "warnings": warnings}
