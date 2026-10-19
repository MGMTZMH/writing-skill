"""Safety tests for docx_utils.py.

Covers three things the skill's programs must never get wrong:
  1. Rewriting a paragraph (proofread): text boxes, links, reference-manager
     fields, content controls, footnote marks and tracked changes are kept intact,
     or the paragraph is left exactly as it was and the call raises ValueError.
  2. Sentence spans: where one sentence ends and the next begins in academic
     prose (abbreviations, initials, decimals, citations, closing quotes).
  3. Sentence-range comments: a comment highlights the sentence that holds a
     phrase, never changes the text, and keeps its markers outside links and fields.

Builds small .docx files in a temporary folder; writes nothing else.

Run from any folder:

    python scripts/test_docx_safety.py

Prints PASS or FAIL for each check (ASCII only); exits 0 when every check passes, 1 otherwise.
"""
import sys

sys.dont_write_bytecode = True

import importlib.util  # noqa: E402
import os  # noqa: E402
import tempfile  # noqa: E402
import zipfile  # noqa: E402
from pathlib import Path  # noqa: E402

from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from lxml import etree  # noqa: E402

SCRIPTS = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location("docx_utils_under_test", str(SCRIPTS / "docx_utils.py"))
du = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(du)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = ('xmlns:w="%s" xmlns:v="urn:schemas-microsoft-com:vml" '
      'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
      'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"' % W)
LQ, RQ = "\u201c", "\u201d"
results = []


def q(tag):
    return "{%s}%s" % (W, tag)


def check(name, ok, detail=""):
    results.append(bool(ok))
    print("%s  %s" % ("PASS" if ok else "FAIL", name.encode("ascii", "backslashreplace").decode()))
    if detail and not ok:
        print("        " + str(detail)[-700:].encode("ascii", "replace").decode())


def refused(p, new_text):
    """True when replace_paragraph_text raises ValueError and leaves the paragraph exactly as it was."""
    before = xml(p._p)
    try:
        du.replace_paragraph_text(p, new_text)
    except ValueError:
        return xml(p._p) == before
    return False


def xml(el):
    return etree.tostring(el)


# ---------------------------------------------------------------------------
# building blocks
# ---------------------------------------------------------------------------
def vml_box(text, floating):
    style = "position:absolute;margin-left:10pt;width:120pt;height:40pt" if floating else "width:120pt;height:40pt"
    return parse_xml('<w:r %s><w:pict><v:shape style="%s"><v:textbox><w:txbxContent><w:p><w:r><w:t>%s</w:t>'
                     '</w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>' % (NS, style, text))


def dml_box(text, floating):
    frame = "wp:anchor" if floating else "wp:inline"
    return parse_xml(
        '<w:r %s><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><%s>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:txbxContent></wps:txbx>'
        '</wps:wsp></a:graphicData></a:graphic></%s></w:drawing></mc:Choice><mc:Fallback><w:pict>'
        '<v:shape style="%s"><v:textbox><w:txbxContent><w:p><w:r><w:t>%s</w:t></w:r></w:p></w:txbxContent>'
        '</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>'
        % (NS, frame, text, frame, "position:absolute" if floating else "", text))


def hyperlink(text):
    return parse_xml('<w:hyperlink %s r:id="rId99"><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:hyperlink>'
                     % (NS, text))


def field_runs(display):
    return [parse_xml('<w:r %s><w:fldChar w:fldCharType="begin"/></w:r>' % NS),
            parse_xml('<w:r %s><w:instrText xml:space="preserve"> ADDIN ZOTERO_ITEM CSL_CITATION {} '
                      '</w:instrText></w:r>' % NS),
            parse_xml('<w:r %s><w:fldChar w:fldCharType="separate"/></w:r>' % NS),
            parse_xml('<w:r %s><w:t>%s</w:t></w:r>' % (NS, display)),
            parse_xml('<w:r %s><w:fldChar w:fldCharType="end"/></w:r>' % NS)]


def sdt(display):
    return parse_xml('<w:sdt %s><w:sdtPr><w:tag w:val="CITATION"/></w:sdtPr><w:sdtContent><w:r><w:t>%s</w:t>'
                     '</w:r></w:sdtContent></w:sdt>' % (NS, display))


def fld_simple(display):
    return parse_xml('<w:fldSimple %s w:instr=" REF _Ref1 \\h "><w:r><w:t>%s</w:t></w:r></w:fldSimple>'
                     % (NS, display))


def footnote_ref(fid):
    return parse_xml('<w:r %s><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
                     '<w:footnoteReference w:id="%s"/></w:r>' % (NS, fid))


def tracked(kind, text):
    tag = "ins" if kind == "ins" else "del"
    inner = "<w:t>%s</w:t>" % text if kind == "ins" else "<w:delText>%s</w:delText>" % text
    return parse_xml('<w:%s %s w:id="77" w:author="A" w:date="2026-09-26T00:00:00Z"><w:r>%s</w:r></w:%s>'
                     % (tag, NS, inner, tag))


def text_run(text, bold=False, italic=False):
    props = ("<w:b/>" if bold else "") + ("<w:i/>" if italic else "")
    return parse_xml('<w:r %s>%s<w:t xml:space="preserve">%s</w:t></w:r>'
                     % (NS, "<w:rPr>%s</w:rPr>" % props if props else "", text))


def paragraph(parts):
    """A fresh paragraph element made of runs and other elements, in order."""
    doc = Document()
    p = doc.add_paragraph()
    for part in parts:
        if isinstance(part, str):
            p._p.append(text_run(part))
        elif isinstance(part, list):
            for el in part:
                p._p.append(el)
        else:
            p._p.append(part)
    return p


def box_texts(p_el):
    return ["".join(t.text or "" for t in box.iter(q("t"))) for box in p_el.iter(q("txbxContent"))]


# ---------------------------------------------------------------------------
# 1. rewriting paragraphs
# ---------------------------------------------------------------------------
def test_rewrite():
    new = "Firms respond by changing their policies."

    p = paragraph(["Firms in order to respond ", vml_box("BOX TEXT", True), "change their policies."])
    du.replace_paragraph_text(p, new)
    check("floating VML text box: paragraph rewritten", du.paragraph_view_text(p._p) == new,
          du.paragraph_view_text(p._p))
    check("floating VML text box: box text kept", box_texts(p._p) == ["BOX TEXT"], box_texts(p._p))

    p = paragraph(["Firms in order to respond ", dml_box("DML BOX", True), "change their policies."])
    du.replace_paragraph_text(p, new)
    check("floating DrawingML text box with fallback: rewritten, both copies of the box text kept",
          du.paragraph_view_text(p._p) == new and box_texts(p._p) == ["DML BOX", "DML BOX"], box_texts(p._p))

    for label, part in (("inline VML text box", vml_box("INLINE", False)),
                        ("inline DrawingML object", dml_box("INLINE", False)),
                        ("content control (reference-manager citation)", sdt("(Smith, 2020)")),
                        ("simple field (cross-reference)", fld_simple("Table 2")),
                        ("tracked insertion", tracked("ins", "inserted ")),
                        ("tracked deletion", tracked("del", "deleted "))):
        p = paragraph(["Firms in order to respond ", part, "change their policies."])
        before = xml(p._p)
        try:
            du.replace_paragraph_text(p, new)
            raised = False
        except ValueError:
            raised = True
        check("%s: not rewritten (ValueError), paragraph unchanged" % label, raised and xml(p._p) == before)

    p = paragraph(["See ", hyperlink("the linked words"), " for in order to learn more."])
    du.replace_paragraph_text(p, "See the linked words to learn more.")
    links = p._p.findall(q("hyperlink"))
    check("hyperlink: text rewritten around the link, link text once and in place",
          du.paragraph_view_text(p._p) == "See the linked words to learn more." and len(links) == 1
          and "".join(t.text or "" for t in links[0].iter(q("t"))) == "the linked words",
          du.paragraph_view_text(p._p))

    p = paragraph(["See ", hyperlink("the linked words"), " for in order to learn more."])
    before = xml(p._p)
    try:
        du.replace_paragraph_text(p, "See the new link to learn more.")
        raised = False
    except ValueError:
        raised = True
    check("hyperlink text changed by the edit: not rewritten, paragraph unchanged", raised and xml(p._p) == before)

    cite = "(Smith, 2020)"
    p = paragraph(["Protest raises risk ", field_runs(cite), " and in order to act firms concede."])
    du.replace_paragraph_text(p, "Protest raises risk %s and to act firms concede." % cite)
    kids = list(p._p)
    kinds = []
    for k in kids:
        fc = k.find(q("fldChar"))
        kinds.append(fc.get(q("fldCharType")) if fc is not None else ("t:" + "".join(t.text or "" for t in k.findall(q("t")))))
    sep, end = kinds.index("separate"), kinds.index("end")
    check("reference-manager field: text rewritten around it, citation still inside the field",
          du.paragraph_view_text(p._p) == "Protest raises risk %s and to act firms concede." % cite
          and "t:" + cite in kinds[sep + 1:end] and p._p.find(".//" + q("instrText")) is not None, kinds)

    p = paragraph(["Protest raises risk ", field_runs(cite), " and in order to act firms concede."])
    before = xml(p._p)
    try:
        du.replace_paragraph_text(p, "Protest raises risk (Smith 2020) and to act firms concede.")
        raised = False
    except ValueError:
        raised = True
    check("citation text changed by the edit: not rewritten, paragraph unchanged", raised and xml(p._p) == before)

    p = paragraph(["Protest raises risk.", footnote_ref(3), " In order to act, firms concede."])
    view = du.paragraph_view_text(p._p)
    check("footnote mark shown to editors as [^3]", view == "Protest raises risk.[^3] In order to act, firms concede.",
          view)
    du.replace_paragraph_text(p, "Protest raises risk.[^3] To act, firms concede.")
    kids = [k for k in p._p if k.tag == q("r")]
    order = ["FN" if k.find(q("footnoteReference")) is not None else "".join(t.text or "" for t in k.findall(q("t")))
             for k in kids]
    check("footnote mark stays after the words it follows",
          du.paragraph_view_text(p._p) == "Protest raises risk.[^3] To act, firms concede."
          and order.index("FN") == 1 and order[0] == "Protest raises risk.", order)

    p = paragraph(["Protest raises risk.", footnote_ref(3), " In order to act, firms concede."])
    before = xml(p._p)
    try:
        du.replace_paragraph_text(p, "Protest raises risk. To act, firms concede.")
        raised = False
    except ValueError:
        raised = True
    check("footnote mark missing from the new text: not rewritten, paragraph unchanged",
          raised and xml(p._p) == before)

    p = paragraph([])
    p._p.append(parse_xml('<w:r %s><w:t>Tabs</w:t><w:tab/><w:t>inside a run in order to align.</w:t></w:r>' % NS))
    before = xml(p._p)
    try:
        du.replace_paragraph_text(p, "Tabs inside a run to align.")
        raised = False
    except ValueError:
        raised = True
    check("result that would not read as intended (tab inside a run): restored, ValueError",
          raised and xml(p._p) == before)

    mixed = parse_xml('<w:r %s><w:t xml:space="preserve">Autonomy in order to </w:t>%s</w:r>'
                      % (NS, etree.tostring(dml_box("M", True).find(q("AlternateContent").replace(
                          W, "http://schemas.openxmlformats.org/markup-compatibility/2006")))
                         .decode()))
    p = paragraph([])
    p._p.append(mixed)
    p._p.append(text_run("matters", italic=True))
    p._p.append(text_run(" for recalls."))
    du.replace_paragraph_text(p, "Autonomy matters for recalls.")
    check("run holding text and a floating object: object not copied by the italics restore",
          len(box_texts(p._p)) == 2 and du.paragraph_view_text(p._p) == "Autonomy matters for recalls.",
          (box_texts(p._p), du.paragraph_view_text(p._p)))

    p = paragraph(["Firms ", tracked("ins", "inserted "), "respond."])
    check("count_tracked_changes counts an insertion", du.count_tracked_changes(p._p) == 1)
    doc = Document()
    body = doc.element.body
    for kind in ("ins", "del"):
        doc.add_paragraph()._p.append(tracked(kind, "x"))
    run = doc.add_paragraph()._p
    run.append(parse_xml('<w:r %s><w:rPr><w:b/><w:rPrChange w:id="5" w:author="A" w:date="2026-09-26T00:00:00Z">'
                         '<w:rPr/></w:rPrChange></w:rPr><w:t>bold now</w:t></w:r>' % NS))
    check("count_tracked_changes counts insertions, deletions and formatting changes",
          du.count_tracked_changes(body) == 3, du.count_tracked_changes(body))
    doc = Document()
    body = doc.element.body
    p = doc.add_paragraph()._p
    p.insert(0, parse_xml('<w:pPr %s><w:pPrChange w:id="6" w:author="A" w:date="2026-09-26T00:00:00Z"><w:pPr/>'
                          '</w:pPrChange></w:pPr>' % NS))
    p.append(parse_xml('<w:moveFrom %s w:id="7" w:author="A" w:date="2026-09-26T00:00:00Z"><w:r><w:delText>x'
                       '</w:delText></w:r></w:moveFrom>' % NS))
    p.append(parse_xml('<w:moveTo %s w:id="8" w:author="A" w:date="2026-09-26T00:00:00Z"><w:r><w:t>x</w:t>'
                       '</w:r></w:moveTo>' % NS))
    check("count_tracked_changes counts moves and paragraph-format changes", du.count_tracked_changes(body) == 3,
          du.count_tracked_changes(body))

    # --- content the rewrite could not place: the paragraph is locked and left as it was ---
    math = parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                     '<m:r><m:t>y</m:t></m:r></m:oMath>')
    check("inline equation: not rewritten, paragraph unchanged",
          refused(paragraph(["where ", math, " is in order to count the recalls."]),
                  "where is to count the recalls."))
    sym = parse_xml('<w:r %s><w:sym w:font="Symbol" w:char="F062"/></w:r>' % NS)
    check("symbol-font character: not rewritten, paragraph unchanged",
          refused(paragraph(["The estimate ", sym, " is in order to show the effect."]),
                  "The estimate is to show the effect."))
    sup = parse_xml('<w:r %s><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:t>2</w:t></w:r>' % NS)
    check("superscript text (R squared): not rewritten, paragraph unchanged",
          refused(paragraph(["The R", sup, " rises in order to 0.42."]), "The R2 rises to 0.42."))
    tag = parse_xml('<w:smartTag %s w:uri="urn:x" w:element="place"><w:r><w:t>Boston</w:t></w:r></w:smartTag>' % NS)
    check("text inside an element the program does not know: not rewritten, paragraph unchanged",
          refused(paragraph(["Plants in ", tag, " act in order to comply."]), "Plants in Boston act to comply."))
    # character formatting that changes what the text shows, when it differs between runs
    def styled(props, text):
        return parse_xml('<w:r %s><w:rPr>%s</w:rPr><w:t xml:space="preserve">%s</w:t></w:r>' % (NS, props, text))
    for label, props, text in (
            ("Greek letter typed in a symbol font", '<w:rFonts w:ascii="Symbol" w:hAnsi="Symbol"/><w:i/>', "b"),
            ("hidden text", "<w:vanish/>", "hidden note "),
            ("struck-through text", "<w:strike/>", "old words "),
            ("small capitals", "<w:smallCaps/>", "Tech "),
            ("capitals set by formatting", "<w:caps/>", "nasa "),
            ("raised text", '<w:position w:val="6"/>', "up ")):
        check("%s among plain runs: not rewritten, paragraph unchanged" % label,
              refused(paragraph(["The estimate ", styled(props, text), "is in order to show the effect."]),
                      "The estimate %sis to show the effect." % text))
    check("text in a different character style (Emphasis) among plain runs: not rewritten, paragraph unchanged",
          refused(paragraph(["The ", styled('<w:rStyle w:val="Emphasis"/>', "key "), "result holds in order to test."]),
                  "The key result holds to test."))
    p = paragraph([styled('<w:rStyle w:val="Emphasis"/>', "All in one style "),
                   styled('<w:rStyle w:val="Emphasis"/>', "in order to test.")])
    du.replace_paragraph_text(p, "All in one style to test.")
    runs = [r for r in p._p.findall(q("r")) if "".join(t.text or "" for t in r.findall(q("t")))]
    check("a paragraph wholly in one character style is rewritten and keeps the style",
          du.paragraph_view_text(p._p) == "All in one style to test."
          and all(r.find(q("rPr")) is not None and r.find(q("rPr")).find(q("rStyle")) is not None
                  and r.find(q("rPr")).find(q("rStyle")).get(q("val")) == "Emphasis" for r in runs),
          du.paragraph_view_text(p._p))
    if os.name == "nt":
        bs = chr(92)
        unc = bs * 2 + "server" + bs + "share" + bs + "x" * 250 + bs + "f.docx"
        got = str(du.long_path(unc))
        check("a long network (UNC) path gets the long-path UNC prefix",
              got.startswith(bs * 2 + "?" + bs + "UNC" + bs + "server" + bs + "share" + bs), got)
    p = paragraph([styled("<w:smallCaps/>", "Every run in small caps "), styled("<w:smallCaps/>", "in order to test.")])
    du.replace_paragraph_text(p, "Every run in small caps to test.")
    runs = [r for r in p._p.findall(q("r")) if r.findall(q("t")) and "".join(t.text or "" for t in r.findall(q("t")))]
    check("a paragraph set wholly in small capitals is rewritten and keeps them",
          du.paragraph_view_text(p._p) == "Every run in small caps to test."
          and all(r.find(q("rPr")) is not None and r.find(q("rPr")).find(q("smallCaps")) is not None for r in runs),
          du.paragraph_view_text(p._p))

    reason = du.paragraph_lock_reason(paragraph(["where ", parse_xml(
        '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>y</m:t>'
        '</m:r></m:oMath>'), " holds."])._p) or ""
    check("the lock reason names the equation", "equation" in reason, reason)

    # --- footnote marks must keep their numbers and order ---
    p = paragraph(["A holds.", footnote_ref(3), " B holds.", footnote_ref(4), " In order to end."])
    check("footnote marks swapped by an edit: not rewritten, paragraph unchanged",
          refused(p, "B holds.[^4] A holds.[^3] To end."))

    # --- a citation whose text shows a straight apostrophe still matches ---
    p = paragraph(["Firms concede ", field_runs("(O'Reilly, 2020)"), " in order to end protests."])
    du.replace_paragraph_text(p, "Firms concede (O'Reilly, 2020) to end protests.")
    check("citation field with a straight apostrophe: text rewritten around it, field kept",
          du.paragraph_view_text(p._p).replace("\u2019", "'") == "Firms concede (O'Reilly, 2020) to end protests."
          and p._p.find(".//" + q("instrText")) is not None, du.paragraph_view_text(p._p))

    # --- an existing comment inside a rewritten paragraph now covers the whole paragraph ---
    p = paragraph(["Firms act. ", parse_xml('<w:commentRangeStart %s w:id="40"/>' % NS), "Teams in order to learn ",
                   parse_xml('<w:commentRangeEnd %s w:id="40"/>' % NS),
                   parse_xml('<w:r %s><w:commentReference w:id="40"/></w:r>' % NS), "adapt."])
    notes = du.replace_paragraph_text(p, "Firms act. Teams learn and adapt.")
    kids = [k.tag for k in p._p]
    first_content = next(i for i, t in enumerate(kids) if t == q("r"))
    check("an existing comment whose words do not survive is widened to the whole paragraph and noted",
          kids.index(q("commentRangeStart")) < first_content
          and kids.index(q("commentRangeEnd")) > max(i for i, k in enumerate(p._p)
                                                      if k.tag == q("r") and k.findall(q("t")))
          and notes and "comment" in notes[0], (kids, notes))

    p = paragraph(["Firms act. ", parse_xml('<w:commentRangeStart %s w:id="41"/>' % NS), "Teams learn fast.",
                   parse_xml('<w:commentRangeEnd %s w:id="41"/>' % NS),
                   parse_xml('<w:r %s><w:commentReference w:id="41"/></w:r>' % NS), " Plants adapt in order to survive."])
    notes = du.replace_paragraph_text(p, "Firms act. Teams learn fast. Plants adapt to survive.")
    covered, inside = "", False
    for el in p._p.iter():
        if el.tag == q("commentRangeStart"):
            inside = True
        elif el.tag == q("commentRangeEnd"):
            inside = False
        elif el.tag == q("t") and inside:
            covered += el.text or ""
    check("an existing comment whose words survive stays on their sentence and is noted",
          du.paragraph_view_text(p._p) == "Firms act. Teams learn fast. Plants adapt to survive."
          and covered == "Teams learn fast." and notes and "sentence" in notes[0], (covered, notes))


# ---------------------------------------------------------------------------
# 2. sentence spans
# ---------------------------------------------------------------------------
SENTENCES = [
    ("Firms respond. Managers act.", ["Firms respond.", "Managers act."]),
    ("Protest raises risk (Harrow & Quill, 2006). Firms concede.",
     ["Protest raises risk (Harrow & Quill, 2006).", "Firms concede."]),
    ("Smith et al. (2020) found X. Y follows.", ["Smith et al. (2020) found X.", "Y follows."]),
    ("The effect is significant (p < .05). The next test fails.",
     ["The effect is significant (p < .05).", "The next test fails."]),
    ("The value is 0.42. It rises.", ["The value is 0.42.", "It rises."]),
    ("J. Smith argued this. Then others agreed.", ["J. Smith argued this.", "Then others agreed."]),
    ("Does autonomy matter? We think so!  Yes.", ["Does autonomy matter?", "We think so!", "Yes."]),
    ("He called it %sthe line.%s Then he left." % (LQ, RQ), ["He called it %sthe line.%s" % (LQ, RQ), "Then he left."]),
    ("Firms (e.g., Apple) act. Teams follow.", ["Firms (e.g., Apple) act.", "Teams follow."]),
    ("See Fig. 2 for details. Then read on.", ["See Fig. 2 for details.", "Then read on."]),
    ("Data from the U.S. Census are used. They cover 2019.", ["Data from the U.S. Census are used.", "They cover 2019."]),
    ("Introduction", ["Introduction"]),
    ("  Leading space. Trailing space.  ", ["Leading space.", "Trailing space."]),
    # known limit: a single capital letter before a period reads as an initial
    ("Appendix A. The next part.", ["Appendix A. The next part."]),
]


def test_sentences():
    for text, expected in SENTENCES:
        got = [text[s:e] for s, e in du.sentence_spans(text)]
        check("sentence spans: %s" % text.strip()[:40], got == expected, got)


# ---------------------------------------------------------------------------
# 3. sentence-range comments
# ---------------------------------------------------------------------------
def ranges(docx_path):
    """{comment id: highlighted text}, and the XML of document.xml."""
    with zipfile.ZipFile(str(docx_path)) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    open_ids, out = set(), {}
    for el in root.iter():
        if el.tag == q("commentRangeStart"):
            open_ids.add(el.get(q("id")))
            out.setdefault(el.get(q("id")), "")
        elif el.tag == q("commentRangeEnd"):
            open_ids.discard(el.get(q("id")))
        elif el.tag == q("t"):
            for cid in open_ids:
                out[cid] += el.text or ""
    return out, root


def comment_ids(docx_path):
    with zipfile.ZipFile(str(docx_path)) as z:
        if "word/comments.xml" not in z.namelist():
            return []
        root = etree.fromstring(z.read("word/comments.xml"))
    return [c.get(q("id")) for c in root.findall(q("comment"))]


def save(tmp, name, parts_list):
    doc = Document()
    for parts in parts_list:
        p = doc.add_paragraph()
        for part in parts:
            if isinstance(part, str):
                p._p.append(text_run(part))
            elif isinstance(part, list):
                for el in part:
                    p._p.append(el)
            else:
                p._p.append(part)
    path = tmp / name
    doc.save(str(path))
    return path


def texts_of(path):
    doc = Document(str(path))
    return [du.paragraph_view_text(p._p) for p in doc.paragraphs]


def test_comments(tmp):
    s1, s2, s3 = "Firms respond to protest. ", "Managers cannot predict targets. ", "The threat spreads."
    path = save(tmp, "one.docx", [["Title line"], [s1, s2 + s3]])
    before = texts_of(path)
    added, notes = du.inject_comments_ex(path, [
        {"search": "Managers cannot predict", "text": "On sentence two.", "author": "T",
         "highlight": "Managers cannot predict"}])
    got, root = ranges(path)
    check("comment highlights exactly the sentence holding the phrase",
          added == 1 and not notes and list(got.values()) == [s2.strip()], (got, notes))
    check("paragraph text unchanged by the comment", texts_of(path) == before)
    ref_ok = False
    for el in root.iter(q("commentRangeEnd")):
        nxt = el.getnext()
        ref_ok = nxt is not None and nxt.find(q("commentReference")) is not None
    check("comment reference sits right after the range end", ref_ok)

    path = save(tmp, "two.docx", [[s1 + s2 + s3]])
    du.inject_comments_ex(path, [{"search": "respond to protest", "text": "Spans two.", "author": "T",
                                  "highlight": "protest. Managers cannot"}])
    got, _ = ranges(path)
    check("a phrase across two sentences highlights both", list(got.values()) == [(s1 + s2).strip()], got)

    path = save(tmp, "three.docx", [[s1 + s2 + s3]])
    n = du.inject_comments(path, [("respond to protest", "Whole paragraph.", "T")])
    got, _ = ranges(path)
    check("3-tuple call keeps the whole-paragraph range and returns a count",
          n == 1 and list(got.values()) == [s1 + s2 + s3], got)

    path = save(tmp, "four.docx", [[s1 + s2 + s3]])
    added, notes = du.inject_comments_ex(path, [{"search": "respond to protest", "text": "x", "author": "T",
                                                 "highlight": "a phrase that is not there"}])
    got, _ = ranges(path)
    check("phrase not found: whole paragraph, and the call reports it",
          added == 1 and notes and list(got.values()) == [s1 + s2 + s3], (got, notes))

    path = save(tmp, "five.docx", [["Firms act. Firms act again. Firms act."]])
    added, notes = du.inject_comments_ex(path, [{"search": "act again", "text": "x", "author": "T",
                                                 "highlight": "Firms act"}])
    got, _ = ranges(path)
    check("phrase occurring twice in the paragraph: whole paragraph, and reported",
          added == 1 and notes and list(got.values()) == ["Firms act. Firms act again. Firms act."], (got, notes))

    path = save(tmp, "six.docx", [["Plain start. ", text_run("Bold middle. Bold end.", bold=True)]])
    du.inject_comments_ex(path, [{"search": "Plain start", "text": "x", "author": "T", "highlight": "Bold end"}])
    got, root = ranges(path)
    bold_runs = [r for r in root.iter(q("r")) if r.find(q("rPr")) is not None and r.find(q("rPr")).find(q("b")) is not None]
    bold_text = "".join("".join(t.text or "" for t in r.findall(q("t"))) for r in bold_runs)
    check("split run keeps its formatting on both sides",
          list(got.values()) == ["Bold end."] and bold_text == "Bold middle. Bold end.", (got, bold_text))

    path = save(tmp, "seven.docx", [["See ", hyperlink("the linked words. More"), " words after the link."]])
    du.inject_comments_ex(path, [{"search": "See", "text": "x", "author": "T", "highlight": "More words after"}])
    got, root = ranges(path)
    link = next(root.iter(q("hyperlink")))
    inside = [c.tag for c in link if c.tag in (q("commentRangeStart"), q("commentRangeEnd"))]
    check("sentence boundary inside a link: markers stay outside the link",
          not inside and "".join(t.text or "" for t in link.iter(q("t"))) == "the linked words. More", (got, inside))

    path = save(tmp, "eight.docx", [["Protest raises risk ", field_runs("(Smith, 2020). Next"), " part follows."]])
    du.inject_comments_ex(path, [{"search": "Protest", "text": "x", "author": "T", "highlight": "Next part"}])
    got, root = ranges(path)
    p = root.find(".//" + q("p"))
    kids = list(p)
    start = [i for i, k in enumerate(kids) if k.tag == q("commentRangeStart")]
    begin = [i for i, k in enumerate(kids) if k.find(q("fldChar")) is not None
             and k.find(q("fldChar")).get(q("fldCharType")) == "begin"]
    check("sentence boundary inside a field result: start marker placed before the field",
          start and begin and start[0] < begin[0], (start, begin))

    path = save(tmp, "eight_b.docx", [["Firms act ", field_runs("(Smith, 2020). Then"), " they stop."]])
    du.inject_comments_ex(path, [{"search": "Firms act", "text": "x", "author": "T", "highlight": "Firms act"}])
    _, root = ranges(path)
    kids = list(root.find(".//" + q("p")))
    end_mark = [i for i, k in enumerate(kids) if k.tag == q("commentRangeEnd")]
    end_run = [i for i, k in enumerate(kids) if k.find(q("fldChar")) is not None
               and k.find(q("fldChar")).get(q("fldCharType")) == "end"]
    check("sentence ending inside a field result: end marker placed after the field",
          end_mark and end_run and end_mark[0] > end_run[0], (end_mark, end_run))

    for label, element in (("content control", sdt("(Smith, 2020). Next")),
                           ("simple field", fld_simple("Table 2. Next"))):
        path = save(tmp, "wrap_%s.docx" % label.replace(" ", "_"),
                    [["Protest raises risk ", element, " part follows."]])
        du.inject_comments_ex(path, [{"search": "Protest", "text": "x", "author": "T", "highlight": "part follows"}])
        _, root = ranges(path)
        kids = list(root.find(".//" + q("p")))
        wrapper = [i for i, k in enumerate(kids) if k.tag in (q("sdt"), q("fldSimple"))]
        start = [i for i, k in enumerate(kids) if k.tag == q("commentRangeStart")]
        inside = [el for el in root.iter(q("commentRangeStart")) if el.getparent().tag != q("p")]
        check("sentence starting inside a %s: start marker placed before it" % label,
              start and wrapper and start[0] < wrapper[0] and not inside, (start, wrapper))

    path = save(tmp, "nine.docx", [[s1 + s2 + s3]])
    du.inject_comments_ex(path, [
        {"search": "respond to protest", "text": "first", "author": "T", "highlight": "Managers cannot"},
        {"search": "respond to protest", "text": "second", "author": "T", "highlight": "cannot predict"}])
    got, _ = ranges(path)
    check("two comments on one sentence: both highlight it, distinct ids",
          sorted(got.values()) == [s2.strip(), s2.strip()] and len(set(got)) == 2, got)

    path = save(tmp, "ten.docx", [[s1 + s2 + s3]])
    du.inject_comments(path, [("respond to protest", "existing", "T")])
    du.inject_comments_ex(path, [{"search": "respond to protest", "text": "new", "author": "T",
                                  "highlight": "The threat"}])
    ids = comment_ids(path)
    check("a draft that already has a comment: old one kept, new id continues",
          ids == ["0", "1"], ids)

    path = save(tmp, "eleven.docx", [["He called it %sthe line,%s  and  then left." % (LQ, RQ)]])
    added, notes = du.inject_comments_ex(path, [{"search": "He called", "text": "x", "author": "T",
                                                 "highlight": 'called it "the line," and then'}])
    check("highlight matched across curly quotes and doubled spaces", added == 1 and not notes, notes)

    path = save(tmp, "eleven_b.docx", [[s1 + s2 + s3]])
    added, notes = du.inject_comments_ex(path, [{"search": "respond to protest", "text": "x", "author": "T",
                                                 "highlight": '"Managers cannot predict"'}])
    got, _ = ranges(path)
    check("a highlight phrase wrapped in quotation marks still matches its sentence",
          added == 1 and not notes and list(got.values()) == [s2.strip()], (got, notes))

    path = save(tmp, "twelve.docx", [["First paragraph."], [s1 + s2 + s3]])
    added, notes = du.inject_comments_ex(path, [{"para_index": 1, "text": "by position", "author": "T",
                                                 "highlight": "The threat"}])
    got, _ = ranges(path)
    check("comment located by paragraph position", added == 1 and list(got.values()) == [s3], (got, notes))

    added, notes = du.inject_comments_ex(path, [{"para_index": 7, "text": "x", "author": "T"}])
    check("paragraph position out of range: not added, reported", added == 0 and notes, notes)


def main():
    test_rewrite()
    test_sentences()
    with tempfile.TemporaryDirectory(prefix="docx_safety_test_") as td:
        test_comments(Path(td))
    print("%d of %d checks passed" % (sum(results), len(results)))
    return 0 if all(results) else 1


if __name__ == "__main__":
    sys.exit(main())
