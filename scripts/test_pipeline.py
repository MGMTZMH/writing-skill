"""End-to-end regression test for the language editor scripts.

Builds small synthetic .docx files in a temporary folder, runs
index_manuscript.py and apply_language_edits.py on them the way a user
would, and checks the results. Nothing outside the temporary folder is
written.

Run from any folder:

    python scripts/test_pipeline.py

Prints PASS or FAIL for each check; exits 0 when every check passes, 1 otherwise.
"""
import datetime
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from lxml import etree

sys.dont_write_bytecode = True

SCRIPTS = Path(__file__).resolve().parent
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
SKILL = SCRIPTS.parent
APPLY = SCRIPTS / "apply_language_edits.py"
INDEX = SCRIPTS / "index_manuscript.py"
ENV_VAR = "LANGUAGE_EDITOR_DOCX_UTILS"
ENV = dict(os.environ, PYTHONDONTWRITEBYTECODE="1", PYTHONIOENCODING="utf-8")
ENV.pop(ENV_VAR, None)
EM_DASH = chr(0x2014)
results = []


def check(name, ok, detail=""):
    results.append(bool(ok))
    print("%s  %s" % ("PASS" if ok else "FAIL", name))
    if detail and not ok:
        tail = str(detail)[-700:].encode("ascii", "replace").decode()
        print("        " + tail.replace("\n", "\n        "))


def run(args, cwd, env=ENV):
    p = subprocess.run([sys.executable] + [str(a) for a in args], cwd=str(cwd), env=env,
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    return p.returncode, p.stdout + p.stderr


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, str(path))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_draft(path, texts):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    doc.save(str(path))
    return path


def scenario(tmp, name, texts):
    """A folder holding one synthetic draft and its paragraph index."""
    folder = tmp / name
    folder.mkdir()
    draft = make_draft(folder / "draft.docx", texts)
    code, out = run([INDEX, "--input", draft, "--out-dir", folder], folder)
    if code != 0:
        raise SystemExit("index_manuscript.py failed:\n" + out)
    return folder, draft


def entry(pid, original, action="REVISE", revised=None, comment="null", importance=None, anchor=None):
    # Test fixtures supply final-text anchors under the tracked-output contract.
    if anchor is None and comment != "null":
        anchor = " ".join((revised if revised is not None else original).split()[:7])
    lines = ["### %s" % pid, "ORIGINAL: %s" % original[:50], "ACTION: %s" % action]
    if revised is not None:
        lines.append("REVISED: %s" % revised)
    lines.append("COMMENT: %s" % comment)
    if importance:
        lines.append("IMPORTANCE: %s" % importance)
    if anchor:
        lines.append("ANCHOR: %s" % anchor)
    return "\n".join(lines + ["", "---", ""])


def comment_ranges(docx_path):
    """{comment text: highlighted text} for every comment in the file."""
    with zipfile.ZipFile(str(docx_path)) as z:
        doc_root = etree.fromstring(z.read("word/document.xml"))
        com_root = etree.fromstring(z.read("word/comments.xml")) if "word/comments.xml" in z.namelist() else None
    w = "{%s}" % W_NS
    texts = {}
    if com_root is not None:
        for c in com_root.findall(w + "comment"):
            texts[c.get(w + "id")] = "".join(t.text or "" for t in c.iter(w + "t"))
    open_ids, got = set(), {}
    for el in doc_root.iter():
        if el.tag == w + "commentRangeStart":
            open_ids.add(el.get(w + "id"))
            got.setdefault(el.get(w + "id"), "")
        elif el.tag == w + "commentRangeEnd":
            open_ids.discard(el.get(w + "id"))
        elif el.tag == w + "t":
            for cid in open_ids:
                got[cid] += el.text or ""
    return {texts.get(cid, cid): rng for cid, rng in got.items()}


def add_xml(paragraph, fragment):
    paragraph._p.append(parse_xml(fragment % ('xmlns:w="%s"' % W_NS)))


def write_edits(path, entries):
    path.write_text("## Edit File\n\n" + "\n".join(entries), encoding="utf-8")
    return path


def write_json(path, data):
    path.write_text(json.dumps(data), encoding="utf-8")
    return path


def reversals(pids):
    return {p: {"class": "reversal", "test": "blind meaning check", "reason": "test"} for p in pids}


def comments_xml(docx_path):
    with zipfile.ZipFile(str(docx_path)) as z:
        return z.read("word/comments.xml").decode("utf-8") if "word/comments.xml" in z.namelist() else ""


def apply_edits(folder, draft, edits, overrides=None, index=True, out_dir="out", apply_script=APPLY):
    destination = folder / out_dir
    counter = 1
    while (destination / "draft_edited_CLEAN.docx").exists():
        destination = folder / (out_dir + "_fresh%d" % counter)
        counter += 1
    out = destination / "draft_edited_CLEAN.docx"
    args = [apply_script, "--input", draft, "--out", out, "--edits", edits]
    if index:
        args += ["--para-index", folder / "para_index.txt"]
    if overrides is not None:
        args += ["--overrides", overrides]
    code, text = run(args, folder)
    manifest_path = out.parent / "language_edit_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    return code, text, manifest, out


def main():
    sys.path.insert(0, str(SCRIPTS))
    du = load_module("docx_utils_under_test", SCRIPTS / "docx_utils.py")
    with tempfile.TemporaryDirectory(prefix="language_editor_test_") as td:
        tmp = Path(td)

        # --- a 21-paragraph draft with 19 revisions, one '--', one flag comment ---
        base = ["Paragraph number %02d says that in order to examine the firms we used a panel." % i
                for i in range(20)]
        base.append('A later paragraph quotes the item "I trust my manager" with straight quotes.')
        folder, draft = scenario(tmp, "basic", base)
        entries = []
        for i in range(20):
            pid = "P%03d" % i
            if i == 5:
                entries.append(entry(pid, base[i], action="KEEP", importance="MINOR",
                                     comment="Long paragraph (9 sentences); a new paragraph could start where the second idea begins."))
            elif i == 2:
                entries.append(entry(pid, base[i], revised="Paragraph number 02 says we examined the firms -- with a panel."))
            else:
                entries.append(entry(pid, base[i], revised="Paragraph number %02d says we examined the firms with a panel." % i))
        write_edits(folder / "language_agent_intro.md", entries)
        empty = write_json(folder / "empty.json", {})

        rev1 = write_json(folder / "rev1.json", reversals(["P003"]))
        code, out, manifest, clean = apply_edits(folder, draft, "language_agent_*.md", rev1, out_dir="o1")
        check("a quoted glob pattern expands and the run succeeds", code == 0 and "parsed 20 entries" in out, out)
        check("a reversal rate of about 5 percent prints no note", code == 0 and "NOTE" not in out, out)
        check("'--' added by an edit is warned about and noted in the manifest",
              "WARN" in out and any("--" in json.dumps(n) for n in manifest.get("notes", [])), out)
        check("a flag comment on a KEEP entry reaches the document", "Long paragraph" in comments_xml(clean), out)
        check("straight quotes the edits never touched give verify errors=0", "verify: errors=0" in out, out)

        brackets = tmp / "run [2026]"
        brackets.mkdir()
        shutil.copyfile(str(folder / "language_agent_intro.md"), str(brackets / "language_agent_intro.md"))
        code, out, _, _ = apply_edits(folder, draft, str(brackets / "language_agent_*.md"), empty, out_dir="o12")
        check("a folder name with square brackets still finds the edit files",
              code == 0 and "parsed 20 entries" in out, out)

        code, out, _, _ = apply_edits(folder, draft, "nothing_*.md", rev1, out_dir="o2")
        check("an edit pattern that matches no file gives a clear message",
              code != 0 and "no edit file matches" in out and "Traceback" not in out, out)

        rev4 = write_json(folder / "rev4.json", reversals(["P003", "P004", "P006", "P007"]))
        code, out, _, _ = apply_edits(folder, draft, folder / "language_agent_intro.md", rev4, out_dir="o3")
        check("a reversal rate of about 21 percent is reported in the summary, with no threshold note",
              code == 0 and "challenged" in out and "NOTE" not in out, out)

        code, out, _, _ = apply_edits(folder, draft, "language_agent_intro.md", folder / "missing.json", out_dir="o4")
        check("a missing overrides file stops the run with a clear message",
              code != 0 and "overrides file not found" in out and "Traceback" not in out, out)

        code, out, manifest, _ = apply_edits(folder, draft, "language_agent_intro.md", empty, out_dir="o5")
        check("an empty overrides file ({}) means no reversals",
              code == 0 and manifest.get("counts", {}).get("reverted") == 0, out)

        code, out, _, clean = apply_edits(folder, draft, "language_agent_intro.md", None, out_dir="o10")
        check("leaving out --overrides stops the run before anything is written",
              code != 0 and "--overrides" in out and "Traceback" not in out and not clean.exists(), out)

        stray = write_json(folder / "stray.json", reversals(["P099"]))
        code, out, manifest, _ = apply_edits(folder, draft, "language_agent_intro.md", stray, out_dir="o11")
        check("an override for a paragraph with no REVISE entry is warned about and noted",
              code == 0 and "matches no REVISE entry" in out
              and any(n.get("pid") == "P099" for n in manifest.get("notes", [])), out)

        # --- span-level partial reversal: one challenged span keeps the other edits ---
        part = write_json(folder / "part.json", {
            "P003": {"class": "reversal", "test": "blind meaning check", "reason": "test"},
            "P002": {"class": "partial", "reason": "the '--' span was challenged; the rest kept",
                     "revised": "Paragraph number 02 says we examined the firms with a panel."}})
        code, out, manifest, clean = apply_edits(folder, draft, "language_agent_intro.md", part, out_dir="o14")
        paras = Document(str(clean)).paragraphs if clean.exists() else []
        refusal = {r.get("pid"): r for r in manifest.get("refusals", [])}
        check("a partial reversal applies the reconstructed text and keeps the paragraph's other edits",
              code == 0 and paras and paras[2].text == "Paragraph number 02 says we examined the firms with a panel."
              and "PARTIAL applied: P002" in out, out)
        check("the challenged spans are logged as a refusal of class partial with both texts",
              refusal.get("P002", {}).get("class") == "partial"
              and refusal.get("P002", {}).get("verdict") == "REVISE -> PARTIAL"
              and "applied_replacement" in refusal.get("P002", {})
              and refusal.get("P003", {}).get("verdict") == "REVISE -> KEEP", manifest)
        check("counts separate full reversals from partials",
              manifest.get("counts", {}).get("reverted") == 1
              and manifest.get("counts", {}).get("partial") == 1
              and manifest.get("counts", {}).get("applied") == 18, manifest.get("counts"))

        part_bad = write_json(folder / "part_bad.json", {"P002": {"class": "partial", "reason": "no text"}})
        code, out, _, _ = apply_edits(folder, draft, "language_agent_intro.md", part_bad, out_dir="o15")
        check("a partial override without reconstructed text stops the run",
              code != 0 and "needs the reconstructed text" in out and "Traceback" not in out, out)

        part_same = write_json(folder / "part_same.json", {
            "P002": {"class": "partial", "reason": "nothing actually dropped",
                     "revised": "Paragraph number 02 says we examined the firms -- with a panel."}})
        code, out, manifest, _ = apply_edits(folder, draft, "language_agent_intro.md", part_same, out_dir="o16")
        check("a partial override identical to the proposal reverses nothing and is noted",
              code == 0 and "matches the proposed revision" in out
              and manifest.get("counts", {}).get("partial") == 0
              and any(n.get("pid") == "P002" for n in manifest.get("notes", [])), out)

        # --- straight quotes introduced by editing ---
        for label, text in (("double", 'An edit that introduced a "straight" quote.'),
                            ("single", "An edit that introduced the editor's straight apostrophe.")):
            doc = Document(str(draft))
            doc.add_paragraph(text)
            doc.save(str(folder / ("added_%s.docx" % label)))
            v = du.verify_docx_formatting(folder / ("added_%s.docx" % label), draft)
            check("a straight %s quote added by editing is an error" % label,
                  any("straight" in e.lower() for e in v["errors"]), v)

        # --- the bundled docx_utils is used from any folder; the override is announced ---
        bundled = str(SCRIPTS / "docx_utils.py").encode("ascii", "replace").decode().lower().replace("/", "\\")
        for label, cwd in (("the temporary folder", tmp), ("the skill folder", SKILL)):
            c, o = run([SCRIPTS / "test_italics_restore.py"], cwd)
            loaded = [line for line in o.splitlines() if line.startswith("docx_utils:")]
            ok = c == 0 and loaded and bundled in loaded[0].lower().replace("/", "\\")
            check("the italics test run from %s loads the bundled docx_utils and passes" % label, ok, o)
        alt = tmp / "alt"
        alt.mkdir()
        shutil.copy(str(SCRIPTS / "docx_utils.py"), str(alt / "docx_utils.py"))
        c, o = run([SCRIPTS / "test_italics_restore.py"], tmp, env=dict(ENV, **{ENV_VAR: str(alt / "docx_utils.py")}))
        check("%s is honored and announced" % ENV_VAR,
              c == 0 and "NOTICE: docx_utils loaded from" in o and str(alt).lower().replace("/", "\\") in o.lower().replace("/", "\\"), o)

        # --- the italics test itself can detect a wrong italic letter ---
        italics = load_module("italics_test_under_test", SCRIPTS / "test_italics_restore.py")
        bad = Document().add_paragraph()
        bad.add_run("This ")
        letter = bad.add_run("p")
        letter.italic = True
        bad.add_run("aper reports p < .05.")
        problems = italics.italic_problems(bad, "p") if hasattr(italics, "italic_problems") else None
        check("the italics test detects one italic letter inside an ordinary word", bool(problems), problems)

        # --- two REVISE entries for one paragraph stop the run before anything is written ---
        dup = tmp / "dup"
        dup.mkdir()
        for name in ("language_agent_a.md", "language_agent_b.md"):
            write_edits(dup / name, [entry("P001", base[1], revised="Paragraph number 01 says %s." % name)])
        code, out, _, clean = apply_edits(folder, draft, str(dup / "language_agent_*.md"), empty, out_dir="o6")
        check("duplicate REVISE entries for one paragraph stop the run",
              code != 0 and "two REVISE entries" in out and not clean.exists(), out)

        # --- a '--' the original already had is not warned about ---
        dash = tmp / "dash"
        dash.mkdir()
        dash_text = "Firms -- in our sample -- grew quickly and then steadily over the whole period."
        dash_draft = make_draft(dash / "draft.docx", [dash_text])
        write_edits(dash / "edits.md", [entry("P000", dash_text, revised="Firms -- in our sample -- grew quickly, then steadily.")])
        code, out, _, _ = apply_edits(dash, dash_draft, dash / "edits.md", write_json(dash / "empty.json", {}), index=False)
        check("a '--' carried over from the original is not warned about",
              code == 0 and "REVISE applied: P000" in out and "contains '--'" not in out, out)

        # --- comments: dashes kept, a [FLAG] tag removed, past-the-end recorded ---
        flags = tmp / "flags"
        flags.mkdir()
        write_edits(flags / "language_agent_flags.md", [
            entry("P006", base[6], action="KEEP", importance="MINOR",
                  comment="[MINOR] A second idea begins here %s consider a split." % EM_DASH),
            entry("P007", base[7], action="KEEP", importance="MAJOR", comment="[FLAG] Consider a transition sentence."),
        ])
        code, out, manifest, clean = apply_edits(folder, draft, flags / "language_agent_flags.md", empty, out_dir="o7")
        cx = comments_xml(clean) if clean.exists() else ""
        check("a dash inside a comment is kept as written", EM_DASH in cx, out)
        check("a leading [FLAG] tag is removed and the comment keeps its words",
              "Consider a transition sentence." in cx and "[FLAG]" not in cx and "MAJOR" not in cx, out)
        check("every retained flag has a sentence anchor and complete coverage",
              manifest.get("counts", {}).get("comments_injected") == 2, manifest)
        write_edits(flags / "missing_target.md", [
            entry("P050", "There is no such paragraph in this short draft at all", action="KEEP", importance="MINOR", comment="Past the end.")])
        missing_code, missing_out, _, missing_clean = apply_edits(folder, draft, flags / "missing_target.md", empty, out_dir="missing_target")
        check("an unresolved comment target blocks publication", missing_code == 5 and not missing_clean.exists(), missing_out)

        # --- a flag reaches Word as the editor wrote it: no importance tag, "Optional:" first ---
        write_edits(flags / "language_agent_plain.md", [
            entry("P010", base[10], action="KEEP", importance="MAJOR",
                  comment="The step from frequent contact to faster fixes is not stated; readers need it for H2."),
            entry("P012", base[12], action="KEEP", importance="MINOR",
                  comment="Optional: this paragraph and the next both run long; keeping the next one short would help."),
            entry("P013", base[13], action="KEEP", importance="MINOR",
                  comment="[MINOR] Say which firms respond. Author decides."),
        ])
        code, out, manifest, clean = apply_edits(folder, draft, flags / "language_agent_plain.md", empty, out_dir="o14")
        got = comment_ranges(clean) if clean.exists() else {}
        check("a plain flag reaches Word exactly as written, with no importance tag",
              code == 0 and "The step from frequent contact to faster fixes is not stated; readers need it for H2." in got,
              (out, list(got)))
        check("an optional flag still opens with 'Optional:' in Word",
              any(k.startswith("Optional: this paragraph and the next") for k in got), list(got))
        check("a leading [MINOR] tag and a closing 'Author decides.' are removed from a flag",
              "Say which firms respond." in got and not any("[M" in k or "Author decides" in k for k in got), list(got))
        imps = {c.get("pid"): c.get("importance") for c in manifest.get("comments", [])}
        check("each flag's importance stays in the manifest",
              imps.get("P010") == "MAJOR" and imps.get("P012") == "MINOR" and imps.get("P013") == "MINOR", imps)

        write_edits(flags / "language_agent_moved.md", [
            entry("P011", base[9], action="KEEP", importance="MINOR", comment="[MINOR] Numbered wrongly on purpose."),
        ])
        code, out, _, clean = apply_edits(folder, draft, flags / "language_agent_moved.md", empty, index=False, out_dir="o12")
        paras = Document(str(clean)).paragraphs if clean.exists() else []
        check("a comment whose number and text disagree goes where the text is, not where the number points",
              code == 0 and paras and "commentRangeStart" in paras[9]._p.xml
              and "commentRangeStart" not in paras[11]._p.xml, out)

        write_edits(flags / "language_agent_bad.md", [entry("P008", base[8], action="MOVE", revised="Moved.")])
        code, out, _, clean = apply_edits(folder, draft, flags / "language_agent_bad.md", empty, out_dir="o8")
        check("an unknown ACTION stops the run before anything is written",
              code != 0 and "unknown ACTION" in out and not clean.exists(), out)

        write_edits(flags / "language_agent_critical.md", [
            entry("P009", base[9], action="KEEP", importance="CRITICAL", comment="[CRITICAL] A level that no longer exists."),
        ])
        code, out, _, clean = apply_edits(folder, draft, flags / "language_agent_critical.md", empty, out_dir="o13")
        check("a CRITICAL label is refused: only MAJOR or MINOR are allowed",
              code != 0 and "MAJOR|MINOR" in out and "Traceback" not in out and not clean.exists(), out)

        # --- the manifest names the rules file that is actually shipped ---
        pkg = tmp / "pkg"
        shutil.copytree(str(SCRIPTS), str(pkg / "scripts"),
                        ignore=shutil.ignore_patterns("__pycache__"))
        (pkg / "references").mkdir()
        rules = sorted((SKILL / "references").glob("rules_language_v*.txt"))
        if rules:
            shutil.copy(str(rules[0]), str(pkg / "references" / "rules_language_v9.txt"))
        code, out, manifest, _ = apply_edits(folder, draft, folder / "language_agent_intro.md", empty, out_dir="o9",
                                             apply_script=pkg / "scripts" / "apply_language_edits.py")
        check("the manifest records the rules file found in references/",
              code == 0 and manifest.get("rule_set") == "rules_language_v9", out)

        # --- paragraph numbers past P999 ---
        big_texts = ["Body paragraph %04d of a long manuscript, written for the index test." % i for i in range(1005)]
        big, big_draft = scenario(tmp, "big", big_texts)
        write_edits(big / "language_agent_big.md", [
            entry("P100", big_texts[100], revised="Body paragraph 0100, revised."),
            entry("P1000", big_texts[1000], revised="Body paragraph 1000, revised."),
        ])
        code, out, _, clean = apply_edits(big, big_draft, big / "language_agent_big.md",
                                          write_json(big / "empty.json", {}))
        paras = Document(str(clean)).paragraphs if clean.exists() else []
        check("P100 and P1000 are both edited, each in its own place",
              code == 0 and len(paras) == 1005 and paras[100].text == "Body paragraph 0100, revised."
              and paras[1000].text == "Body paragraph 1000, revised.", out)

        # --- identical paragraphs: the numbered position decides; never guess ---
        same = "Identical paragraph text that appears twice in this draft for the match test."
        amb, amb_draft = scenario(tmp, "amb", ["Opening paragraph of the ambiguity draft, which is unique.", same,
                                               "A unique middle paragraph that separates the identical ones.", same])
        amb_empty = write_json(amb / "empty.json", {})
        write_edits(amb / "second.md", [entry("P003", same, revised="Identical paragraph text, revised in its second place.")])
        code, out, _, clean = apply_edits(amb, amb_draft, amb / "second.md", amb_empty, out_dir="o1")
        paras = Document(str(clean)).paragraphs if clean.exists() else []
        check("an edit to the second of two identical paragraphs changes the second one",
              code == 0 and paras and paras[3].text.startswith("Identical paragraph text, revised") and paras[1].text == same, out)
        write_edits(amb / "wrong_place.md", [entry("P002", same, revised="Identical paragraph text, revised somewhere.")])
        code, out, manifest, clean = apply_edits(amb, amb_draft, amb / "wrong_place.md", amb_empty, index=False, out_dir="o2")
        paras = Document(str(clean)).paragraphs if clean.exists() else []
        check("identical text away from the numbered position blocks ambiguous editing",
              code == 5 and "ambiguous" in out and not clean.exists(), out)

        # --- verification messages use the P### numbering of para_index.txt ---
        src, edited = Document(), Document()
        for i in range(5):
            src.add_paragraph("Plain paragraph %d for the message format test." % i)
            edited.add_paragraph("Plain paragraph %d for the message format test." % i)
        mixed = src.add_paragraph()
        mixed.add_run("Participants.").bold = True
        mixed.add_run(" We surveyed 402 managers in 2020.")
        edited.add_paragraph("Participants. We surveyed 402 managers in 2020.")
        src.save(str(tmp / "fmt_src.docx"))
        edited.save(str(tmp / "fmt_out.docx"))
        v = du.verify_docx_formatting(tmp / "fmt_out.docx", tmp / "fmt_src.docx")
        check("verification messages name paragraphs as P005, not P5",
              any(e.startswith("P005") for e in v["errors"]), v)

        # --- comments are dated today unless a date is given ---
        dated = tmp / "dated.docx"
        shutil.copy(str(draft), str(dated))
        du.inject_comments(dated, [("Paragraph number 10", "Date test comment.", "Tester")])
        check("an injected comment carries today's date by default",
              datetime.date.today().isoformat() in comments_xml(dated), comments_xml(dated)[-300:])

        # --- unresolved tracked changes stop the run before any work ---
        trk = tmp / "tracked"
        trk.mkdir()
        doc = Document()
        doc.add_paragraph("An opening paragraph without any tracked change in it.")
        p = doc.add_paragraph("Firms ")
        add_xml(p, '<w:ins %s w:id="9" w:author="A" w:date="2026-09-26T00:00:00Z"><w:r><w:t>quietly </w:t></w:r></w:ins>')
        p.add_run("respond in order to protest.")
        trk_draft = trk / "draft.docx"
        doc.save(str(trk_draft))
        code, out = run([INDEX, "--input", trk_draft, "--out-dir", trk / "index_out"], trk)
        check("indexing a draft with tracked changes stops (exit 3) and writes nothing",
              code == 3 and "accept or reject" in out.lower() and not (trk / "index_out").exists(), out)
        write_edits(trk / "edits.md", [entry("P001", "Firms respond in order to protest.",
                                             revised="Firms respond to protest.")])
        code, out, _, clean = apply_edits(trk, trk_draft, trk / "edits.md", write_json(trk / "empty.json", {}),
                                          index=False)
        check("applying edits to a draft with tracked changes stops (exit 3) and writes nothing",
              code == 3 and "accept or reject" in out.lower() and not clean.exists(), out)

        # --- a paragraph holding a content control is locked: listed, never rewritten ---
        lock = tmp / "locked"
        lock.mkdir()
        doc = Document()
        doc.add_paragraph("An opening paragraph that the test leaves alone entirely.")
        p = doc.add_paragraph("Protest raises risk ")
        add_xml(p, '<w:sdt %s><w:sdtPr><w:tag w:val="CITATION"/></w:sdtPr><w:sdtContent><w:r>'
                   '<w:t>(Smith, 2020)</w:t></w:r></w:sdtContent></w:sdt>')
        p.add_run(" and in order to act firms concede.")
        doc.add_paragraph("A plain paragraph that says in order to act firms concede.")
        lock_draft = lock / "draft.docx"
        doc.save(str(lock_draft))
        code, out = run([INDEX, "--input", lock_draft, "--out-dir", lock], lock)
        locked_file = lock / "locked_paragraphs.txt"
        locked_text = locked_file.read_text(encoding="utf-8") if locked_file.exists() else ""
        check("the index lists a paragraph holding a content control as locked",
              code == 0 and "P001" in locked_text and "content control" in locked_text, out + locked_text)
        write_edits(lock / "edits.md", [
            entry("P001", "Protest raises risk", revised="Protest raises risk and to act firms concede."),
            entry("P002", "A plain paragraph that says in order to act firms concede.",
                  revised="A plain paragraph that says firms concede to act.")])
        code, out, manifest, clean = apply_edits(lock, lock_draft, lock / "edits.md", write_json(lock / "empty.json", {}))
        paras = Document(str(clean)).paragraphs if clean.exists() else []
        skipped = {s.get("pid"): s.get("why", "") for s in manifest.get("skipped", [])}
        check("an edit to a locked paragraph is skipped with the reason; other edits apply",
              code == 0 and "content control" in skipped.get("P001", "")
              and paras and "(Smith, 2020)" in paras[1]._p.xml and "in order to act" in "".join(
                  t.text or "" for t in paras[1]._p.iter("{%s}t" % W_NS))
              and paras[2].text == "A plain paragraph that says firms concede to act.", (skipped, out))

        # --- footnote marks: shown as [^N], kept after their words, or the edit is skipped ---
        fn = tmp / "footnotes"
        fn.mkdir()
        doc = Document()
        doc.add_paragraph("An opening paragraph that carries no footnote mark at all.")
        for words in ("Protest raises risk.", "Boycotts cost sales."):
            p = doc.add_paragraph(words)
            add_xml(p, '<w:r %s><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteReference w:id="2"/></w:r>')
            p.add_run(" In order to act, firms concede.")
        fn_draft = fn / "draft.docx"
        doc.save(str(fn_draft))
        code, out = run([INDEX, "--input", fn_draft, "--out-dir", fn], fn)
        index_text = (fn / "para_index.txt").read_text(encoding="utf-8") if (fn / "para_index.txt").exists() else ""
        check("the paragraph index shows a footnote mark as [^2] where it sits",
              "Protest raises risk.[^2] In order to act" in index_text, index_text)
        write_edits(fn / "edits.md", [
            entry("P001", "Protest raises risk.[^2] In order to act, firms concede.",
                  revised="Protest raises risk.[^2] To act, firms concede."),
            entry("P002", "Boycotts cost sales.[^2] In order to act, firms concede.",
                  revised="Boycotts cost sales. To act, firms concede.")])
        code, out, manifest, clean = apply_edits(fn, fn_draft, fn / "edits.md", write_json(fn / "empty.json", {}))
        paras = Document(str(clean)).paragraphs if clean.exists() else []
        order = []
        if paras:
            for r in paras[1]._p.findall("{%s}r" % W_NS):
                if r.find("{%s}footnoteReference" % W_NS) is not None:
                    order.append("FN")
                else:
                    order.append("".join(t.text or "" for t in r.findall("{%s}t" % W_NS)))
        skipped = {s.get("pid"): s.get("why", "") for s in manifest.get("skipped", [])}
        check("an edit that keeps the [^N] mark is applied and the mark stays after its words",
              code == 0 and order[:2] == ["Protest raises risk.", "FN"]
              and "".join(o for o in order if o != "FN") == "Protest raises risk. To act, firms concede.", order)
        check("an edit that drops the [^N] mark is skipped with the reason",
              "footnote" in skipped.get("P002", ""), skipped)

        # --- note texts for the editors to read; an abstract with a run-in label is found ---
        from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part
        nt = tmp / "notetexts"
        nt.mkdir()
        doc = Document()
        doc.add_paragraph("Abstract. We ask whether teams that hold decision autonomy on the line produce fewer "
                          "recalls, and two studies test the argument across plants.")
        p = doc.add_paragraph("We coded autonomy from plant handbooks.")
        add_xml(p, '<w:r %s><w:footnoteReference w:id="1"/></w:r>')
        notes_xml = ('<w:footnotes xmlns:w="%s">'
                     '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                     '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/>'
                     '</w:r></w:p></w:footnote>'
                     '<w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> Two coders '
                     'read every handbook; agreement was 0.91.</w:t></w:r></w:p></w:footnote></w:footnotes>' % W_NS)
        doc.part.relate_to(Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, notes_xml.encode("utf-8"),
                                doc.part.package), RT.FOOTNOTES)
        nt_draft = nt / "draft.docx"
        doc.save(str(nt_draft))
        code, out = run([INDEX, "--input", nt_draft, "--out-dir", nt], nt)
        notes = (nt / "notes.txt").read_text(encoding="utf-8") if (nt / "notes.txt").exists() else ""
        check("notes.txt gives each note's text by paragraph and mark, for reading only",
              code == 0 and "P001 [^1] footnote: Two coders read every handbook; agreement was 0.91." in notes,
              (out, notes))
        sections = (nt / "sections.txt").read_text(encoding="utf-8") if (nt / "sections.txt").exists() else ""
        check("an abstract that opens with a run-in label is found, however long the paragraph",
              "PRESENT  abstract" in sections, sections)

        # --- flags point to a sentence; drift notes become comments; overrides are checked ---
        anc_texts = ["An opening paragraph that stays exactly as it is in this test.",
                     "Teams decide fast. Plants with autonomy recall fewer products. Managers resist change.",
                     "Firms respond in order to protest. The threat spreads to peers.",
                     "A paragraph whose flag phrase is wrong. It has two sentences.",
                     "Drift paragraph in order to test notes. Second sentence stays.",
                     "Partial paragraph in order to test notes, in order to keep. Tail stays."]
        anc, anc_draft = scenario(tmp, "anchors", anc_texts)
        write_edits(anc / "language_agent_anchors.md", [
            entry("P001", anc_texts[1], action="KEEP", importance="MINOR",
                  comment="Autonomy is used before it is defined.", anchor="autonomy recall fewer"),
            entry("P002", anc_texts[2], revised="Firms respond to protest. The threat spreads to peers.",
                  importance="MAJOR", comment="Name the peers.", anchor="spreads to peers"),
            entry("P002", anc_texts[2], action="KEEP", importance="MINOR",
                  comment="Say which firms respond.", anchor="Firms respond to protest"),
            entry("P003", anc_texts[3], action="KEEP", importance="MINOR",
                  comment="Clarify this flag phrase.", anchor="flag phrase is wrong"),
            entry("P004", anc_texts[4], revised="Drift paragraph to test notes. Second sentence stays."),
            entry("P005", anc_texts[5], revised="Partial paragraph to test notes, to keep. Tail stays.")])
        drift = write_json(anc / "drift.json", {
            "P004": {"class": "drift", "notes": [{"anchor": "to test notes", "note": "narrows the scope slightly"}]},
            "P005": {"class": "partial", "reason": "the second span was challenged",
                     "revised": "Partial paragraph to test notes, in order to keep. Tail stays.",
                     "notes": [{"anchor": "to test notes", "note": "a slight shift in emphasis"}]}})
        code, out, manifest, clean = apply_edits(anc, anc_draft, anc / "language_agent_anchors.md", drift)
        got = comment_ranges(clean) if clean.exists() else {}
        def rng(prefix):
            return next((v for k, v in got.items() if k.startswith(prefix)), None)
        check("a flag on a KEEP paragraph highlights the sentence its ANCHOR names",
              code == 0 and rng("Autonomy is used") == "Plants with autonomy recall fewer products.", (got, out))
        check("a flag on a revised paragraph highlights the revised sentence",
              rng("Name the peers") == "The threat spreads to peers.", got)
        check("a second flag on a revised paragraph, in its own KEEP entry, highlights its own sentence",
              rng("Say which firms") == "Firms respond to protest."
              and rng("Name the peers") == "The threat spreads to peers.", got)
        notes_text = json.dumps(manifest.get("notes", []))
        check("a retained flag is anchored to its actual final sentence",
              rng("Clarify this flag phrase") == "A paragraph whose flag phrase is wrong.", got)
        paras = Document(str(clean)).paragraphs if clean.exists() else []
        check("a drift override keeps the edit and adds a check comment on its sentence",
              paras and paras[4].text == "Drift paragraph to test notes. Second sentence stays."
              and rng("Please check this sentence") == "Drift paragraph to test notes.", (got, out))
        drift_pids = [d.get("pid") for d in manifest.get("drift", [])]
        counts = manifest.get("counts", {})
        check("drift notes are listed in the manifest and do not count as reversals",
              sorted(drift_pids) == ["P004", "P005"] and counts.get("reverted") == 0
              and counts.get("partial") == 1 and counts.get("drift_notes") == 2, manifest)
        check("the manifest counts the comments listed and the comments placed",
              counts.get("comments_listed") == counts.get("comments_injected") == len(got), counts)
        check("a partial override may carry drift notes too",
              paras and paras[5].text == "Partial paragraph to test notes, in order to keep. Tail stays."
              and sum(1 for k in got if k.startswith("Please check this sentence")) == 2, got)
        lost = write_json(anc / "drift_lost_anchor.json", {
            "P004": {"class": "drift", "notes": [{"anchor": "words this paragraph lacks", "note": "reads broader"}]}})
        code, out, man8, _ = apply_edits(anc, anc_draft, anc / "language_agent_anchors.md", lost, out_dir="o8")
        drift_notes = [n for n in man8.get("notes", []) if n.get("pid") == "P004"]
        check("a missing drift anchor blocks publication",
              code == 4 and not (anc / "o8" / "draft_edited_CLEAN.docx").exists(), out)
        no_anchor = write_json(anc / "drift_no_anchor.json", {
            "P004": {"class": "drift", "notes": [{"note": "reads broader"}]}})
        code, out, _, clean9 = apply_edits(anc, anc_draft, anc / "language_agent_anchors.md", no_anchor, out_dir="o9")
        check("a drift note without an anchor stops the run before anything is written",
              code == 1 and "anchor" in out and "Traceback" not in out and not clean9.exists(), out)
        # --- exit codes: a stop before anything is written is 1; a failure while adding comments is 4 ---
        failing = anc / "failing_docx_utils.py"
        failing.write_text(
            "import importlib.util, sys\n"
            "_spec = importlib.util.spec_from_file_location('real_docx_utils', %r)\n"
            "_real = importlib.util.module_from_spec(_spec)\n"
            "sys.modules['real_docx_utils'] = _real\n"
            "_spec.loader.exec_module(_real)\n"
            "globals().update({k: v for k, v in vars(_real).items() if not k.startswith('__')})\n"
            "def _find_phrase(*args, **kwargs):\n"
            "    return None, 0\n" % str(SCRIPTS / "docx_utils.py"),
            encoding="utf-8")
        out_c = anc / "o4" / "draft_edited_CLEAN.docx"
        code, out = run([APPLY, "--input", anc_draft, "--out", out_c, "--edits", anc / "language_agent_anchors.md",
                         "--para-index", anc / "para_index.txt", "--overrides", drift], anc,
                        env=dict(ENV, LANGUAGE_EDITOR_DOCX_UTILS=str(failing)))
        mpath = anc / "o4" / "language_edit_manifest.json"
        mnotes = json.dumps(json.loads(mpath.read_text(encoding="utf-8")).get("notes", [])) if mpath.exists() else ""
        check("a comment placement failure exits 4 and publishes no output",
              code == 4 and not out_c.exists() and not mpath.exists(), out)
        code, out = run([INDEX, "--input", anc_draft, "--out-dir", anc / "o6"], anc,
                        env=dict(ENV, LANGUAGE_EDITOR_DOCX_UTILS=str(failing)))
        check("the index program, which is otherwise quiet about it, says when it loads another docx_utils",
              code == 0 and "NOTICE: docx_utils loaded from" in out, out)
        code, out = run([APPLY, "--input", anc_draft, "--out", anc / "o5" / "x.docx",
                         "--edits", anc / "language_agent_anchors.md", "--para-index", anc / "para_index.txt",
                         "--overrides", anc / "no_such_file.json"], anc)
        check("a stop before anything is written exits 1", code == 1 and not (anc / "o5" / "x.docx").exists(), out)

        bad_drift = write_json(anc / "bad_drift.json", {"P004": {"class": "drift"}})
        code, out, _, clean = apply_edits(anc, anc_draft, anc / "language_agent_anchors.md", bad_drift, out_dir="o2")
        check("a drift override without notes stops the run before anything is written",
              code != 0 and "notes" in out and "Traceback" not in out and not clean.exists(), out)
        odd = write_json(anc / "odd.json", {"P004": {"class": "revers", "reason": "typo in the class"}})
        code, out, _, clean = apply_edits(anc, anc_draft, anc / "language_agent_anchors.md", odd, out_dir="o3")
        check("an unknown override class stops the run before anything is written",
              code != 0 and "class" in out and "Traceback" not in out and not clean.exists(), out)

        # --- Windows paths over 260 characters (long-path support switched off) ---
        if os.name == "nt":
            deep = du.long_path(tmp / ("deep_folder_" + "d" * 50) / ("nested_folder_" + "e" * 50) / ("f" * 40))
            deep.mkdir(parents=True)
            stem = "a_draft_with_a_long_file_name_for_the_path_test"
            make_draft(du.long_path(deep / (stem + ".docx")),
                       ["A paragraph that says in order to examine the firms we used a panel."])
            plain = str(deep)[4:] if str(deep).startswith("\\\\?\\") else str(deep)
            out_dir = os.path.join(plain, "language_editor_output", stem)
            code, out = run([INDEX, "--input", os.path.join(plain, stem + ".docx"), "--out-dir", out_dir], tmp)
            out_long = du.long_path(out_dir)
            wrote_index = (out_long / "para_index.txt").exists()
            if wrote_index:
                write_edits(out_long / "language_agent_long.md", [
                    entry("P000", "A paragraph that says in order to examine the firms we used a panel.",
                          revised="A paragraph that says we used a panel to examine the firms.")])
                write_json(out_long / "verifier_reverts.json", {})
            clean_name = stem + "_edited_CLEAN.docx"
            code2, out2 = run([APPLY, "--input", os.path.join(plain, stem + ".docx"),
                               "--out", os.path.join(out_dir, clean_name),
                               "--edits", os.path.join(out_dir, "language_agent_*.md"),
                               "--para-index", os.path.join(out_dir, "para_index.txt"),
                               "--overrides", os.path.join(out_dir, "verifier_reverts.json")], tmp)
            check("index and apply work when the output paths exceed 260 characters",
                  len(os.path.join(out_dir, clean_name)) > 260 and code == 0 and wrote_index and code2 == 0
                  and (out_long / clean_name).exists(), out + out2)
            # the temporary folder's own cleanup cannot remove paths this long: delete the
            # deep tree through the long-path prefix (forced, since its root is short)
            prefix = chr(92) * 2 + "?" + chr(92)
            shutil.rmtree(prefix + os.path.normpath(str(tmp / ("deep_folder_" + "d" * 50))))

            # a folder just under the prefix threshold, holding an edit file whose path passes 260
            base = os.path.normpath(str(tmp))
            near = os.path.join(base, "n" * (236 - len(base) - 1))
            os.makedirs(near)
            make_draft(du.long_path(os.path.join(near, "d.docx")),
                       ["A paragraph that says in order to examine the firms we used a panel."])
            code, out = run([INDEX, "--input", os.path.join(near, "d.docx"), "--out-dir", near], tmp)
            long_name = "language_agent_" + "x" * 30 + ".md"
            write_edits(du.long_path(os.path.join(near, long_name)), [
                entry("P000", "A paragraph that says in order to examine the firms we used a panel.",
                      revised="A paragraph that says we used a panel to examine the firms.")])
            write_json(du.long_path(os.path.join(near, "verifier_reverts.json")), {})
            code2, out2 = run([APPLY, "--input", os.path.join(near, "d.docx"),
                               "--out", os.path.join(near, "c.docx"),
                               "--edits", os.path.join(near, "language_agent_*.md"),
                               "--para-index", os.path.join(near, "para_index.txt"),
                               "--overrides", os.path.join(near, "verifier_reverts.json")], tmp)
            check("edit files over 260 characters in a folder under the prefix threshold are read",
                  len(near) <= 240 and len(os.path.join(near, long_name)) > 260 and code == 0 and code2 == 0
                  and du.long_path(os.path.join(near, "c.docx")).exists(), out + out2)
            shutil.rmtree(prefix + near)

            # the index's own files, and the manifest next to the output, in folders just under the threshold
            edge = os.path.join(base, "m" * (239 - len(base) - 1))
            os.makedirs(edge)
            make_draft(du.long_path(os.path.join(edge, "d.docx")),
                       ["A paragraph that says in order to examine the firms we used a panel."])
            code, out = run([INDEX, "--input", os.path.join(edge, "d.docx"), "--out-dir", edge], tmp)
            check("the index writes its files in a folder just under the prefix threshold",
                  code == 0 and len(os.path.join(edge, "locked_paragraphs.txt")) > 260
                  and du.long_path(os.path.join(edge, "locked_paragraphs.txt")).exists(), out)
            mf = os.path.join(base, "q" * (233 - len(base) - 1))
            os.makedirs(mf)
            write_edits(du.long_path(os.path.join(edge, "language_agent_e.md")), [
                entry("P000", "A paragraph that says in order to examine the firms we used a panel.",
                      revised="A paragraph that says we used a panel to examine the firms.")])
            write_json(du.long_path(os.path.join(edge, "verifier_reverts.json")), {})
            code2, out2 = run([APPLY, "--input", os.path.join(edge, "d.docx"),
                               "--out", os.path.join(mf, "c.docx"),
                               "--edits", os.path.join(edge, "language_agent_e.md"),
                               "--para-index", os.path.join(edge, "para_index.txt"),
                               "--overrides", os.path.join(edge, "verifier_reverts.json")], tmp)
            check("the manifest is written next to an output file just under the prefix threshold",
                  code2 == 0 and len(os.path.join(mf, "language_edit_manifest.json")) > 260
                  and du.long_path(os.path.join(mf, "language_edit_manifest.json")).exists(), out2)
            shutil.rmtree(prefix + edge)
            shutil.rmtree(prefix + mf)

    print("%d of %d checks passed" % (sum(results), len(results)))
    return 0 if all(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
